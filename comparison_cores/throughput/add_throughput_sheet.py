#!/usr/bin/env python3
"""Add measured comparison-core throughput to the master workbook.

External cores: this harness (out/throughput.csv).  RV-IM100 RV32 cores:
DMIPS/MHz from the master sheet (the software's own cycle-timed window, the same
method as the harness), CoreMark/MHz from revision_0922/T3 (true value,
iterations*f/cycles, not the integer-second figure the board printed).
Area and Fmax: the Core comparison sheet (same core-only methodology for all).
"""
import csv, pathlib, re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

ROOT = pathlib.Path(__file__).resolve().parents[2]
DST  = pathlib.Path.home()/"Documents"/"Research_data_MASTER_0921.xlsx"
TP   = pathlib.Path(__file__).resolve().parent/"out"/"throughput.csv"
T3   = ROOT/"revision_0922"/"T3"/"T3_fragment_corrected_coremark.csv"
F, HDR, WARN = "NanumGothic", PatternFill("solid", fgColor="FFBFBFBF"), PatternFill("solid", fgColor="FFFFF2CC")
SHEET = "Core throughput"

def put(ws, r, c, v, *, bold=False, fill=None, align="left", wrap=False, fmt=None):
    x = ws.cell(row=r, column=c, value=v); x.font = Font(name=F, size=11, bold=bold)
    if fill: x.fill = fill
    x.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fmt: x.number_format = fmt
    return x

wb = openpyxl.load_workbook(DST)
cc = wb["Core comparison"]
area = {}
for r in range(1, cc.max_row + 1):
    n = cc.cell(r, 2).value          # B: core, D: LUT, K: Fmax
    if isinstance(n, str) and isinstance(cc.cell(r, 4).value, (int, float)):
        area[n.strip()] = dict(lut=cc.cell(r, 4).value, fmax=cc.cell(r, 11).value)
rv = wb["RV32 SoC and FPGA"]
dmips = {}
for r in range(1, rv.max_row + 1):
    n, v = rv.cell(r, 2).value, rv.cell(r, 5).value
    if isinstance(n, str) and isinstance(v, str) and "@" in v:
        dmips["RV32" + n.strip()] = float(v.split("@")[0])
t3 = {row["variant"]: row for row in csv.DictReader(open(T3))}
tp = {(r["config"], r["benchmark"]): r for r in csv.DictReader(open(TP))}
assert all(r["status"] == "VALID" for r in tp.values())

rows = []
ext = [("PicoRV32", "picorv32_nola", "picorv32", "RV32IM", "registered-ready memory (dhrystone/testbench_nola.v)"),
       ("PicoRV32", "picorv32_la",   "picorv32", "RV32IM", "look-ahead memory (dhrystone/testbench.v) -- best case"),
       ("VexRiscv", "vexriscv",      "vexriscv_nodebug", "RV32IM", "GenFullNoMmuNoCache, DebugPlugin removed"),
       ("RVCoreP",  "rvcorep",       "rvcorep",  "RV32I",  "v0.5.3 default configuration")]
for core, cfg, akey, isa, note in ext:
    d, c = tp[(cfg, "dhrystone")], tp[(cfg, "coremark")]
    rows.append(dict(core=core, isa=isa, config=note, src="measured, this harness",
                     dmhz=float(d["value"]), cmhz=float(c["value"]), dcpi=float(d["cpi"]), ccpi=float(c["cpi"]),
                     lut=area[akey]["lut"], fmax=area[akey]["fmax"]))
for v in ["RV32I_5SP", "RV32IM_5SP", "RV32IM_6SP", "RV32IM_7SP", "RV32IM_7SP_BRAM",
          "RV32IM_7SP_BRAM_Opt", "RV32IM_8SP_withoutOpt", "RV32IM_8SP"]:
    t = t3[v]
    rows.append(dict(core=v, isa="RV32I" if v.startswith("RV32I_") else "RV32IM", config="RV-IM100",
                     src="DMIPS: master sheet; CoreMark: T3 true value",
                     dmhz=dmips[v], cmhz=float(t["TRUE CoreMark/MHz"]),
                     dcpi=float(t["Dhrystone cycles/iter"]) / float(t["Dhrystone instr/iter"]),
                     ccpi=float(t["CoreMark CPI"]), lut=area[v]["lut"], fmax=area[v]["fmax"]))

if SHEET in wb.sheetnames: wb.remove(wb[SHEET])
ws = wb.create_sheet(SHEET, index=wb.sheetnames.index("Core comparison") + 1)
widths = [24, 9, 44, 11, 13, 10, 10, 9, 11, 13, 14, 34]
for i, w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
r = 1
put(ws, r, 1, "Throughput of the comparison cores, measured under the RV-IM100 methodology", bold=True); r += 2
for t in ["Every core runs the same Dhrystone 2.1 and CoreMark images: RV-IM100's own ports, gcc 15.2.0, "
          "Dhrystone -O2 (300,000 runs), CoreMark -O2 -fno-common -funroll-loops.  rv32im images for PicoRV32 "
          "and VexRiscv, rv32i for RVCoreP.  Only the BSP timer differs: a memory-mapped cycle counter in the "
          "testbench, because RVCoreP has no CSRs.",
          "Cycle-accurate Verilator simulation of each core at the boundary we synthesized, with zero-wait memory "
          "matching each core's own reference model.  DMIPS/MHz = runs*1e6/(cycles*1757); "
          "CoreMark/MHz = iterations*1e6/cycles.  All 8 runs pass Dhrystone's 22 self-checks or CoreMark's CRC validation.",
          "Throughput at Fmax is derived: per-MHz figure x core-only Fmax from the Core comparison sheet.  "
          "Core-only Fmax is not a SoC operating frequency, for any core."]:
    put(ws, r, 1, t, wrap=True); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
    ws.row_dimensions[r].height = 44; r += 1
r += 1
hdr = ["core", "ISA", "configuration", "DMIPS/MHz", "CoreMark/MHz", "CPI (Dhry)", "CPI (CM)", "LUT",
       "Fmax (MHz)", "DMIPS @ Fmax", "CoreMark @ Fmax", "source"]
for i, h in enumerate(hdr): put(ws, r, i + 1, h, bold=True, fill=HDR, align="center" if 0 < i < 11 else "left", wrap=True)
ws.row_dimensions[r].height = 32; r += 1
first = r
for x in rows:
    if x["config"] == "RV-IM100" and x["core"] == "RV32I_5SP":
        r += 1
    vals = [x["core"], x["isa"], x["config"], x["dmhz"], x["cmhz"], x["dcpi"], x["ccpi"], x["lut"], x["fmax"],
            x["dmhz"] * x["fmax"], x["cmhz"] * x["fmax"], x["src"]]
    fmts = [None, None, None, "0.000", "0.000", "0.000", "0.000", "0", "0.0", "0.0", "0.0", None]
    for i, (v, f) in enumerate(zip(vals, fmts)):
        put(ws, r, i + 1, v, align="center" if 0 < i < 11 and i != 2 else "left", fmt=f, wrap=(i == 2))
    r += 1
r += 1
put(ws, r, 1, "Read these before citing", bold=True, fill=WARN)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12); r += 1
for t in ["Absolute DMIPS/MHz here must not be set beside published DMIPS/MHz from other papers.  This port is built -ffreestanding with a byte-wise strcpy and retires 406 instructions per Dhrystone run; letting gcc expand the string copy (builtins + newlib) cuts it to 267, and in that form PicoRV32 reproduces its published 0.516 (0.515) and VexRiscv its published 1.21 (1.235).  Optimisation level changes it by one instruction.  The same applies to RV-IM100's own DMIPS/MHz.",
          "PicoRV32: its README notes the look-ahead interface is usually not usable at maximum clock, so the "
          "registered-ready row is the one to pair with Fmax; the look-ahead row is a best case.",
          "RVCoreP is RV32I.  Compare it with RV32I_5SP.  Its CoreMark/MHz is low because without M it executes "
          "712,323 instructions per iteration against 279,947 for rv32im.",
          "RVCoreP IPC here (Dhrystone 0.944, CoreMark 0.807) agrees with its paper (0.935, 0.823) despite a "
          "different compiler and benchmark build -- a check that the harness does not add stalls.",
          "RV-IM100 CoreMark/MHz is the true value from revision_0922/T3.  The RV32/RV64 SoC and FPGA sheets still "
          "carry the board's integer-second figures, which are 0.1-7.3% high (and 9.75% low for RV32IM_8SP_withoutOpt)."]:
    put(ws, r, 1, t, wrap=True); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
    ws.row_dimensions[r].height = 46; r += 1

toc = wb["Contents"]
cr = next(i for i in range(1, toc.max_row + 1) if toc.cell(i, 2).value == "Core comparison")
if toc.cell(cr + 1, 2).value != SHEET:
    toc.insert_rows(cr + 1, 1)
put(toc, cr + 1, 2, SHEET, bold=True)
put(toc, cr + 1, 3, "DMIPS/MHz, CoreMark/MHz and CPI of PicoRV32, VexRiscv and RVCoreP, measured on RV-IM100's own "
                    "benchmark images in cycle-accurate simulation; RV-IM100 RV32 cores alongside.", wrap=True)
toc.row_dimensions[cr + 1].height = 28
for i in range(1, toc.max_row + 1):
    v = toc.cell(i, 3).value
    if isinstance(v, str) and v.startswith("External soft-core CoreMark/MHz and DMIPS/MHz are not measured"):
        put(toc, i, 3, "External-core throughput is measured in simulation (Core throughput sheet), not on the board.  "
                       "Absolute DMIPS/MHz is specific to this Dhrystone build and must not be compared with "
                       "published figures from other papers.", wrap=True)
wb.save(DST)
print(f"  wrote '{SHEET}' ({len(rows)} rows) to {DST.name}")
for x in rows:
    print(f"  {x['core']:<22}{x['isa']:<8}{x['dmhz']:>7.3f}{x['cmhz']:>8.3f}{x['dcpi']:>8.3f}{x['ccpi']:>8.3f}"
          f"{x['lut']:>6}{x['fmax']:>9.1f}{x['dmhz']*x['fmax']:>8.1f}{x['cmhz']*x['fmax']:>8.1f}  {x['config'][:40]}")
