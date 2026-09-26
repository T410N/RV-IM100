#!/usr/bin/env python3
"""T14: emit every number the paper uses, from data, never by hand.

    paper_numbers.py [--figures]      -> revision_0922/paper_numbers.csv
                                         (+ revision_0922/figures/*.pdf with --figures)

Inputs
  * Research_data_MASTER_0921.xlsx for everything the revision does not change;
  * the revision fragments for everything it corrects:
      T3  true CoreMark it/s and per-iteration cycles/instructions (CoreMark and Dhrystone)
      T4  authoritative SAIF power and normalised metrics
      T5  huffbench/slre/wikisort re-measured with a 12 KB stack
      T1  RV32 Embench FPGA (board runs pending; simulation reference used, labelled)

Every output row: key, value, unit, source, note.  `source` names the sheet!cell or
file the value comes from; derived rows name their inputs.  Projections are labelled.
"""
import csv, math, os, re, sys
import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
REV = f'{ROOT}/revision_0922'
MASTER = f'{ROOT}/Research_data_MASTER_0921.xlsx'
OUT = f'{REV}/paper_numbers.csv'

VAR = ['I_5SP', 'IM_5SP', 'IM_6SP', 'IM_7SP', 'IM_7SP_BRAM', 'IM_7SP_BRAM_Opt', 'IM_8SP_withoutOpt', 'IM_8SP']
LABEL = {'I_5SP': 'I-5', 'IM_5SP': 'IM-5', 'IM_6SP': 'IM-6', 'IM_7SP': 'IM-7 (IO only)', 'IM_7SP_BRAM': 'IM-7',
         'IM_7SP_BRAM_Opt': 'IM-7 Opt', 'IM_8SP_withoutOpt': 'IM-8 noOpt', 'IM_8SP': 'IM-8'}
# (name, from, to): the paper's main path plus the ablation steps
TRANSITIONS = [('M extension', 'I_5SP', 'IM_5SP'), ('5->6', 'IM_5SP', 'IM_6SP'),
               ('6->7 (paper, IO+BRAM)', 'IM_6SP', 'IM_7SP_BRAM'), ('6->7 IO stage only', 'IM_6SP', 'IM_7SP'),
               ('7 IO -> 7 BRAM', 'IM_7SP', 'IM_7SP_BRAM'), ('7 -> 7 Opt', 'IM_7SP_BRAM', 'IM_7SP_BRAM_Opt'),
               ('7 -> 8 noOpt (EXR only)', 'IM_7SP_BRAM', 'IM_8SP_withoutOpt'), ('8 noOpt -> 8', 'IM_8SP_withoutOpt', 'IM_8SP'),
               ('7 Opt -> 8 (EXR on optimised)', 'IM_7SP_BRAM_Opt', 'IM_8SP'), ('7 -> 8 (paper)', 'IM_7SP_BRAM', 'IM_8SP'),
               ('5 -> 8 overall', 'IM_5SP', 'IM_8SP')]
DMIPS = 1757.0

rows = []


def emit(key, value, unit, source, note=''):
    if isinstance(value, float):
        value = round(value, 6)
    rows.append((key, value, unit, source, note))


def f(x):
    return float(str(x).strip())


def at_mhz(text):
    """'1.105 @ 45MHz' -> (1.105, 45.0)"""
    m = re.match(r'\s*([0-9.]+)\s*@\s*([0-9.]+)\s*MHz', str(text))
    return float(m.group(1)), float(m.group(2))


def rcsv(p):
    return list(csv.DictReader(open(p, newline=''), delimiter='\t' if p.endswith('.tsv') else ','))


def geomean(xs):
    return math.exp(sum(math.log(x) for x in xs) / len(xs))


def load():
    wb = openpyxl.load_workbook(MASTER, data_only=True)
    D = {}
    for arch in ('RV32', 'RV64'):
        ws = wb[f'{arch} SoC and FPGA']
        for r in range(3, 11):
            lab = ws.cell(r, 2).value
            if lab not in VAR:
                continue
            v = D.setdefault((arch, lab), {})
            v['dhry_per_s'] = (f(ws.cell(r, 3).value), f'{ws.title}!C{r}')
            v['cm_reported'] = (f(ws.cell(r, 4).value), f'{ws.title}!D{r}')
            v['dhry_mhz'] = (at_mhz(ws.cell(r, 5).value)[1], f'{ws.title}!E{r}')
            v['cm_mhz'] = (at_mhz(ws.cell(r, 8).value)[1], f'{ws.title}!H{r}')
            v['core_fmax'] = (f(ws.cell(r, 10).value), f'{ws.title}!J{r}')
            v['core_lut'] = (f(ws.cell(r, 20).value), f'{ws.title}!T{r}')
            v['core_ff'] = (f(ws.cell(r, 21).value), f'{ws.title}!U{r}')
            v['core_dsp'] = (f(ws.cell(r, 22).value), f'{ws.title}!V{r}')
        block = None
        for r in range(11, ws.max_row + 1):
            lab = ws.cell(r, 2).value
            if lab in ('SoC Dhrystone', 'SoC Coremark'):
                block = 'dhry' if 'Dhry' in lab else 'cm'
                continue
            if lab == 'Legend':
                break
            if block and lab in VAR:
                v = D[(arch, lab)]
                v[f'soc_fmax_{block}'] = (f(ws.cell(r, 10).value), f'{ws.title}!J{r}')
                v[f'soc_lut_{block}'] = (f(ws.cell(r, 3).value), f'{ws.title}!C{r}')
                v[f'soc_bram_{block}'] = (f(ws.cell(r, 6).value), f'{ws.title}!F{r}')
    # T3 corrections
    t3 = f'{REV}/T3/T3_fragment_corrected_coremark.csv'
    for r in rcsv(t3):
        arch, lab = r['variant'][:4], r['variant'][4:]
        v = D[(arch, lab)]
        src = f'T3/T3_fragment_corrected_coremark.csv[{r["variant"]}]'
        v['cm_true'] = (f(r['TRUE it/s = f / cycles-per-iter']), src)
        v['cm_cyc_iter'] = (f(r['CoreMark cycles/iter (sim)']), src)
        v['cm_ins_iter'] = (f(r['CoreMark instr/iter']), src)
        v['dh_cyc_iter'] = (f(r['Dhrystone cycles/iter']), src)
        v['dh_ins_iter'] = (f(r['Dhrystone instr/iter']), src)
    # T4
    for r in rcsv(f'{REV}/T4/T4_normalized_metrics.csv'):
        arch, lab = r['variant'][:4], r['variant'][4:]
        D[(arch, lab)]['norm'] = (r, f'T4/T4_normalized_metrics.csv[{r["variant"]}]')
    for r in rcsv(f'{REV}/T4/T4_saif_authoritative.csv'):
        arch, lab = r['variant'][:4], r['variant'][4:]
        D[(arch, lab)][f'saif_{r["bench"]}'] = (r, f'T4/T4_saif_authoritative.csv[{r["variant"]},{r["bench"]}]')
    return wb, D


def main():
    wb, D = load()
    val = lambda a, v, k: D[(a, v)][k][0]
    src = lambda a, v, k: D[(a, v)][k][1]

    # ---- 1. frequency and throughput per variant
    for a in ('RV32', 'RV64'):
        for v in VAR:
            p = f'{a}.{v}'
            for k, unit in (('core_fmax', 'MHz'), ('soc_fmax_dhry', 'MHz'), ('soc_fmax_cm', 'MHz'),
                            ('dhry_mhz', 'MHz'), ('cm_mhz', 'MHz')):
                emit(f'{p}.{k}', val(a, v, k), unit, src(a, v, k),
                     {'core_fmax': 'post-synthesis, core only, 5 ns over-constraint',
                      'soc_fmax_dhry': 'post-route, Dhrystone image, at the build PLL constraint',
                      'soc_fmax_cm': 'post-route, CoreMark image, at the build PLL constraint',
                      'dhry_mhz': 'run clock of the Dhrystone board run', 'cm_mhz': 'run clock of the CoreMark board run'}[k])
            dps = val(a, v, 'dhry_per_s')
            emit(f'{p}.dhrystones_per_s', dps, '1/s', src(a, v, 'dhry_per_s'))
            emit(f'{p}.dmips', dps / DMIPS, 'DMIPS', src(a, v, 'dhry_per_s'), 'Dhrystones/s / 1757')
            emit(f'{p}.dmips_per_mhz', dps / DMIPS / val(a, v, 'dhry_mhz'), 'DMIPS/MHz',
                 f"{src(a, v, 'dhry_per_s')}, {src(a, v, 'dhry_mhz')}")
            emit(f'{p}.coremark', val(a, v, 'cm_true'), 'it/s', src(a, v, 'cm_true'),
                 'true score f*iterations/cycles; master value biased by integer-second truncation')
            emit(f'{p}.coremark_reported_master', val(a, v, 'cm_reported'), 'it/s', src(a, v, 'cm_reported'), 'superseded')
            emit(f'{p}.coremark_per_mhz', val(a, v, 'cm_true') / val(a, v, 'cm_mhz'), 'CM/MHz',
                 f"{src(a, v, 'cm_true')}, {src(a, v, 'cm_mhz')}")
            emit(f'{p}.cpi_dhrystone', val(a, v, 'dh_cyc_iter') / val(a, v, 'dh_ins_iter'), 'cycles/instr', src(a, v, 'dh_cyc_iter'),
                 'per-iteration, RTL simulation (== FPGA)')
            emit(f'{p}.cpi_coremark', val(a, v, 'cm_cyc_iter') / val(a, v, 'cm_ins_iter'), 'cycles/instr', src(a, v, 'cm_cyc_iter'),
                 'per-iteration, RTL simulation (== FPGA)')
            emit(f'{p}.dhrystone_cycles_per_iter', val(a, v, 'dh_cyc_iter'), 'cycles', src(a, v, 'dh_cyc_iter'))
            emit(f'{p}.coremark_cycles_per_iter', val(a, v, 'cm_cyc_iter'), 'cycles', src(a, v, 'cm_cyc_iter'))
            for k in ('core_lut', 'core_ff', 'core_dsp'):
                emit(f'{p}.{k}', val(a, v, k), 'count', src(a, v, k), 'core-only synthesis, memory excluded')
            for k in ('soc_lut_dhry', 'soc_bram_dhry', 'soc_lut_cm', 'soc_bram_cm'):
                emit(f'{p}.{k}', val(a, v, k), 'count', src(a, v, k), 'SoC, includes memories')
            # implementation variance: same RTL, two images
            emit(f'{p}.soc_fmax_image_spread', abs(val(a, v, 'soc_fmax_dhry') - val(a, v, 'soc_fmax_cm')), 'MHz',
                 f"{src(a, v, 'soc_fmax_dhry')} vs {src(a, v, 'soc_fmax_cm')}", 'same RTL, Dhrystone vs CoreMark image')
            # normalised
            n, ns = D[(a, v)]['norm']
            for k, unit in (('CM_per_W', 'CM/W'), ('mJ_per_CM_iter', 'mJ'), ('CM_per_W_core_attr', 'CM/W'),
                            ('mJ_per_CM_iter_core_attr', 'mJ'), ('uJ_per_Dhry_iter', 'uJ'), ('uJ_per_Dhry_iter_core_attr', 'uJ'),
                            ('CM_per_kLUT_core', 'CM/kLUT')):
                if n.get(k):
                    emit(f'{p}.{k}', f(n[k]), unit, ns, 'core_attr = dynamic - PLL - I/O' if 'core_attr' in k else '')
            emit(f'{p}.dmips_per_klut_core', dps / DMIPS / val(a, v, 'core_lut') * 1000, 'DMIPS/kLUT',
                 f"{src(a, v, 'dhry_per_s')}, {src(a, v, 'core_lut')}")
            for b in ('dhrystone', 'coremark'):
                s, ss = D[(a, v)][f'saif_{b}']
                emit(f'{p}.power_{b}_total', f(s['saif_total']), 'W', ss, f"SAIF, {s['source']}")
                emit(f'{p}.power_{b}_dynamic', f(s['saif_dynamic']), 'W', ss)
                emit(f'{p}.power_{b}_core_attr', f(s['core_attributable_dyn']), 'W', ss, 'dynamic - PLL - I/O')
                emit(f'{p}.power_{b}_saif_coverage', f(s['coverage_pct']), '%', ss, 'read_saif nets matched')

    # ---- 2. per-transition deltas and break-even
    for a in ('RV32', 'RV64'):
        for name, x, y in TRANSITIONS:
            p = f'{a}.transition[{name}]'
            ins = f'{a}.{x} -> {a}.{y}'
            d = lambda k: val(a, y, k) / val(a, x, k) - 1
            for k, lab in (('core_fmax', 'core_fmax'), ('soc_fmax_dhry', 'soc_fmax_dhry'), ('dhry_mhz', 'run_mhz_dhry'),
                           ('dhry_per_s', 'dhrystones_per_s'), ('cm_true', 'coremark'), ('core_lut', 'core_lut'),
                           ('core_ff', 'core_ff')):
                emit(f'{p}.delta_{lab}', 100 * d(k), '%', ins)
            cpi_x = val(a, x, 'dh_cyc_iter') / val(a, x, 'dh_ins_iter')
            cpi_y = val(a, y, 'dh_cyc_iter') / val(a, y, 'dh_ins_iter')
            emit(f'{p}.delta_cpi_dhrystone', 100 * (cpi_y / cpi_x - 1), '%', ins, 'per-iteration CPI')
            cm_x = val(a, x, 'cm_cyc_iter') / val(a, x, 'cm_ins_iter')
            cm_y = val(a, y, 'cm_cyc_iter') / val(a, y, 'cm_ins_iter')
            emit(f'{p}.delta_cpi_coremark', 100 * (cm_y / cm_x - 1), '%', ins, 'per-iteration CPI')
            # break-even: throughput changes by f_ratio / cycles-per-iteration ratio
            fr = val(a, y, 'dhry_mhz') / val(a, x, 'dhry_mhz')
            cr = val(a, y, 'dh_cyc_iter') / val(a, x, 'dh_cyc_iter')
            emit(f'{p}.breakeven_dhry_freq_ratio', fr, 'x', ins, 'run-clock ratio')
            emit(f'{p}.breakeven_dhry_cycles_ratio', cr, 'x', ins, 'cycles-per-iteration ratio (includes instr-count change)')
            emit(f'{p}.breakeven_dhry_predicted_change', 100 * (fr / cr - 1), '%', ins, 'freq ratio / cycles ratio')
            emit(f'{p}.breakeven_dhry_measured_change', 100 * d('dhry_per_s'), '%', ins, 'board Dhrystones/s')
            fr = val(a, y, 'cm_mhz') / val(a, x, 'cm_mhz')
            cr = val(a, y, 'cm_cyc_iter') / val(a, x, 'cm_cyc_iter')
            emit(f'{p}.breakeven_cm_predicted_change', 100 * (fr / cr - 1), '%', ins, 'freq ratio / cycles ratio')
            emit(f'{p}.breakeven_cm_measured_change', 100 * d('cm_true'), '%', ins, 'true CoreMark')

    # ---- 3. 2x2 ablation around the 7->8 step
    for a in ('RV32', 'RV64'):
        base, opt, exr, both = 'IM_7SP_BRAM', 'IM_7SP_BRAM_Opt', 'IM_8SP_withoutOpt', 'IM_8SP'
        for k, lab in (('dhry_per_s', 'dhrystones_per_s'), ('cm_true', 'coremark'), ('dhry_mhz', 'run_mhz'),
                       ('soc_fmax_dhry', 'soc_fmax'), ('core_fmax', 'core_fmax')):
            e = lambda v: val(a, v, k) / val(a, base, k) - 1
            eo, ee, eb = e(opt), e(exr), e(both)
            p = f'{a}.ablation.{lab}'
            s = f'{a}.{{{base},{opt},{exr},{both}}}.{k}'
            emit(f'{p}.optimisations_only', 100 * eo, '%', s, 'vs 7SP_BRAM')
            emit(f'{p}.exr_only', 100 * ee, '%', s, 'vs 7SP_BRAM')
            emit(f'{p}.both', 100 * eb, '%', s, 'vs 7SP_BRAM')
            emit(f'{p}.sum_of_individual', 100 * (eo + ee), '%', s)
            emit(f'{p}.product_of_individual', 100 * ((1 + eo) * (1 + ee) - 1), '%', s, 'multiplicative expectation')
            emit(f'{p}.subadditivity', 100 * (eb - (eo + ee)), 'pp', s, 'both - (opt + exr); negative = sub-additive')
        for bench, ck, ik in (('dhrystone', 'dh_cyc_iter', 'dh_ins_iter'), ('coremark', 'cm_cyc_iter', 'cm_ins_iter')):
            cpi = lambda v: val(a, v, ck) / val(a, v, ik)
            for v, lab in ((opt, 'optimisations_only'), (exr, 'exr_only'), (both, 'both')):
                emit(f'{a}.ablation.cpi_{bench}.{lab}', 100 * (cpi(v) / cpi(base) - 1), '%', f'{a}.{v} vs {a}.{base}')

    # ---- 4. cross-width ratios at every depth
    for v in VAR:
        for k, lab in (('core_lut', 'core_lut'), ('core_ff', 'core_ff'), ('core_dsp', 'core_dsp'), ('core_fmax', 'core_fmax'),
                       ('soc_fmax_dhry', 'soc_fmax_dhry')):
            r32, r64 = val('RV32', v, k), val('RV64', v, k)
            if r32:
                emit(f'crosswidth.{v}.{lab}_rv64_over_rv32', r64 / r32, 'x', f'RV64.{v} / RV32.{v}')
        d32 = val('RV32', v, 'dhry_per_s') / DMIPS / val('RV32', v, 'dhry_mhz')
        d64 = val('RV64', v, 'dhry_per_s') / DMIPS / val('RV64', v, 'dhry_mhz')
        emit(f'crosswidth.{v}.dmips_per_mhz_rv64_over_rv32', d64 / d32, 'x', f'RV64.{v} / RV32.{v}')
        c32 = val('RV32', v, 'cm_true') / val('RV32', v, 'cm_mhz')
        c64 = val('RV64', v, 'cm_true') / val('RV64', v, 'cm_mhz')
        emit(f'crosswidth.{v}.coremark_per_mhz_rv64_over_rv32', c64 / c32, 'x', f'RV64.{v} / RV32.{v}')

    # ---- 5. Embench: FPGA runtimes (RV64 measured; RV32 pending) and projections
    fpga = {}
    ws = wb['Embench FPGA RV64']
    bench = None
    for r in range(1, ws.max_row + 1):
        h = ws.cell(r, 2).value
        if h and str(h).startswith('RV64'):
            bench = str(h).split()[-1]
            continue
        if bench and h in VAR and ws.cell(r, 3).value not in (None, ''):
            fpga[('RV64', h, bench)] = (f(ws.cell(r, 3).value), f(ws.cell(r, 9).value), f'Embench FPGA RV64!C{r}/I{r}')
    t1 = f'{ROOT}/riscof-env/logs/embench_rv32_fpga_vs_sim.csv'
    rv32_src = 'measured' if os.path.exists(t1) else 'simulation (FPGA pending)'
    clk32 = {'I_5SP': 45, 'IM_5SP': 43, 'IM_6SP': 50, 'IM_7SP': 57.501, 'IM_7SP_BRAM': 71.999, 'IM_7SP_BRAM_Opt': 92,
             'IM_8SP_withoutOpt': 113.999, 'IM_8SP': 122}
    if os.path.exists(t1):
        for r in rcsv(t1):
            if r.get('fpga_cycles'):
                fpga[('RV32', r['variant'][4:], r['benchmark'])] = (f(r['fpga_cycles']), clk32[r['variant'][4:]],
                                                                    'FPGA: ' + os.path.relpath(t1, ROOT))
    else:
        for r in rcsv(f'{REV}/T1/embench_rv32_simref_v4.csv'):
            fpga[('RV32', r['variant'][4:], r['benchmark'])] = (f(r['sim_cycles']), clk32[r['variant'][4:]], 'T1/embench_rv32_simref_v4.csv (sim; FPGA pending)')
    benches = ['matmult-int', 'crc32', 'nettle-aes', 'statemate', 'md5sum']
    for a in ('RV32', 'RV64'):
        for v in VAR:
            ratios = []
            for b in benches:
                if (a, v, b) not in fpga:
                    continue
                cyc, mhz, s = fpga[(a, v, b)]
                ms = cyc / mhz / 1000
                emit(f'{a}.{v}.embench_fpga.{b}.runtime_ms', ms, 'ms', s,
                     ('FPGA' if a == 'RV64' else rv32_src) + f', {mhz} MHz build clock')
                ref = fpga.get((a, 'IM_5SP', b))
                if ref:
                    rr = ms / (ref[0] / ref[1] / 1000)
                    ratios.append(rr)
                    emit(f'{a}.{v}.embench_fpga.{b}.runtime_vs_IM_5SP', rr, 'x', s)
            if len(ratios) == len(benches):
                emit(f'{a}.{v}.embench_fpga.geomean_runtime_vs_IM_5SP', geomean(ratios), 'x', 'geomean of the 5 ratios',
                     'RV64 IM_5SP nettle-aes build is 30 MHz (did not close at 38)' if a == 'RV64' else rv32_src)
    # projections for all 19 simulated benchmarks at each variant's Embench build clock
    ws = wb['Embench simulation']
    cols = {ws.cell(5, c).value: c for c in range(3, ws.max_column + 1)}
    over = {(r['variant'], r['benchmark']): f(r['cycles']) for r in rcsv(f'{REV}/T5/T5_fragment_embench_12KBstack.csv')}
    clk64 = {v: fpga[('RV64', v, 'matmult-int')][1] for v in VAR if ('RV64', v, 'matmult-int') in fpga}
    for r in range(6, ws.max_row + 1):
        full = ws.cell(r, 2).value
        if not full or not str(full).startswith('RV'):
            continue
        a, v = full[:4], full[4:]
        mhz = (clk32 if a == 'RV32' else clk64).get(v)
        for b, c in cols.items():
            cyc = over.get((full, b), ws.cell(r, c).value)
            srcc = 'T5 12 KB stack' if (full, b) in over else f'Embench simulation!{openpyxl.utils.get_column_letter(c)}{r}'
            if isinstance(cyc, (int, float)) and mhz:
                emit(f'{a}.{v}.embench_projected.{b}.runtime_ms', cyc / mhz / 1000, 'ms', srcc,
                     f'PROJECTION: simulated cycles at the {mhz} MHz Embench build clock')

    with open(OUT, 'w', newline='') as fh:
        w = csv.writer(fh)
        w.writerow(['key', 'value', 'unit', 'source', 'note'])
        w.writerows(rows)
    print(f'{len(rows)} numbers -> {OUT}')
    if '--figures' in sys.argv:
        import paper_figures
        paper_figures.make(D, VAR, LABEL, fpga, benches, f'{REV}/figures')


if __name__ == '__main__':
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    main()
