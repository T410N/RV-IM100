#!/usr/bin/env python3
"""Collect every revision fragment into one review workbook.

    revision_0922/Revision_review_0923.xlsx

Sheet 'Change list' names each cell of Research_data_MASTER_0921.xlsx that a
fragment changes, with the current value, the proposed value, the reason and the
evidence file.  The other sheets carry the fragments themselves; derived
quantities are live formulas over their inputs, so the reviewer can see how
each number is formed.  Nothing in the master file is modified.
"""
import csv, json, os, re, glob, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
REV = f'{ROOT}/revision_0922'
ENV = f'{ROOT}/riscof-env'
MASTER = f'{ROOT}/Research_data_MASTER_0921.xlsx'
OUT = f'{REV}/Revision_review_0923.xlsx'

F = 'Arial'
HDR_FILL = PatternFill('solid', fgColor='FFD9D9D9')
CHG_FILL = PatternFill('solid', fgColor='FFFFF2CC')     # proposed change
BAD_FILL = PatternFill('solid', fgColor='FFF8CBAD')     # invalid / contradicts
THIN = Side(style='thin', color='FFBFBFBF')


def rcsv(path, header=True, names=None):
    with open(path, newline='') as f:
        if header:
            return list(csv.DictReader(f, delimiter='\t' if path.endswith('.tsv') else ','))
        return [dict(zip(names, r)) for r in csv.reader(f)]


def num(x):
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return x
    s = str(x).strip()
    if s in ('', 'None'):
        return None
    if s in ('True', 'False'):
        return 'yes' if s == 'True' else 'no'      # plain text, not an =TRUE() formula
    try:
        return int(s)
    except ValueError:
        try:
            return float(s)
        except ValueError:
            return s


def title(ws, text, sub=None):
    ws['A1'] = text
    ws['A1'].font = Font(name=F, bold=True, size=13)
    if sub:
        ws['A2'] = sub
        ws['A2'].font = Font(name=F, italic=True, size=9, color='FF595959')
        ws['A2'].alignment = Alignment(wrap_text=False)


def table(ws, row, headers, rows, widths=None, number_formats=None):
    """Write a header row at `row` and data below; returns (first_data_row, last_row)."""
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(name=F, bold=True, size=9)
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = Border(bottom=THIN)
    for i, r in enumerate(rows, row + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=F, size=9)
            if number_formats and headers[j - 1] in number_formats:
                c.number_format = number_formats[headers[j - 1]]
    ws.freeze_panes = ws.cell(row=row + 1, column=2)
    for j, h in enumerate(headers, 1):
        w = (widths or {}).get(h)
        if w is None:
            vals = [len(str(h))] + [len(str(r[j - 1])) for r in rows[:200] if r[j - 1] is not None and not str(r[j - 1]).startswith('=')]
            w = min(max(vals) + 2, 60)
        ws.column_dimensions[get_column_letter(j)].width = w
    return row + 1, row + len(rows)


def csv_sheet(wb, name, path, sub, header=True, names=None, formats=None):
    ws = wb.create_sheet(name)
    title(ws, name, sub + f'   Source: {os.path.relpath(path, ROOT)}')
    rows = rcsv(path, header, names)
    hdr = list(rows[0].keys())
    table(ws, 4, hdr, [[num(r[h]) for h in hdr] for r in rows], number_formats=formats)
    return ws


def main():
    mwb = openpyxl.load_workbook(MASTER, data_only=True)
    wb = openpyxl.Workbook()
    idx = wb.active
    idx.title = 'Index'
    changes = []          # (task, sheet, cell, row label, field, current, proposed, reason, evidence)

    # ---------------------------------------------------------------- data
    cm = rcsv(f'{REV}/T3/T3_fragment_corrected_coremark.csv')
    cmv = {r['variant']: r for r in cm}
    saif = rcsv(f'{REV}/T4/T4_saif_authoritative.csv')
    saifv = {(r['variant'], r['bench']): r for r in saif}
    nm = {r['variant']: r for r in rcsv(f'{REV}/T4/T4_normalized_metrics.csv')}
    audit = rcsv(f'{REV}/T3/soc_rows_from_vivado_reports.csv')
    e12 = rcsv(f'{REV}/T5/T5_fragment_embench_12KBstack.csv')

    # ---------------------------------------------------------------- change list
    # 1. CoreMarks/MHz in the core block of the SoC sheets (column H, rows 3-10)
    for arch in ('RV32', 'RV64'):
        ws = mwb[f'{arch} SoC and FPGA']
        for r in range(3, 11):
            lab = ws.cell(r, 2).value
            if not lab:
                continue
            v = arch + lab
            cur = ws.cell(r, 8).value
            t = cmv[v]
            prop = f"{float(t['TRUE CoreMark/MHz']):.3f} @ {float(t['CoreMark run MHz']):g}MHz"
            why = ('mislabelled 100 MHz image run at 114 MHz, plus integer-second truncation'
                   if v == 'RV32IM_8SP_withoutOpt' else 'integer-second truncation in CoreMark (HAS_FLOAT 0)')
            changes.append(('T3', f'{arch} SoC and FPGA', f'H{r}', lab, 'Coremarks/MHz', cur, prop,
                            f"{why}; master error {float(t['master error %']):+.2f}%",
                            'T3/T3_coremark_anomaly_and_soc_rows.md'))
    # 2. PLL cell RV64 IM_8SP (both SoC blocks), and SAIF columns of every SoC row
    for arch in ('RV32', 'RV64'):
        ws = mwb[f'{arch} SoC and FPGA']
        block = None
        for r in range(1, ws.max_row + 1):
            lab = ws.cell(r, 2).value
            if lab in ('SoC Dhrystone', 'SoC Coremark'):
                block = 'dhrystone' if 'Dhry' in lab else 'coremark'
                continue
            if lab == 'Legend':
                block = None
            if not block or not lab or not re.match(r'I_|IM_', str(lab)):
                continue
            v = arch + lab
            if v == 'RV64IM_8SP' and ws.cell(r, 9).value == 0:
                changes.append(('T7', ws.title, f'I{r}', f'{lab} ({block})', 'PLL', 0, 1,
                                'utilization.rpt: PLLE2_ADV = 1 (all 32 builds use one PLL)',
                                'T7/T7_metadata_checks.md'))
            s = saifv[(v, block)]
            for col, fld, key in ((24, 'SAIF dynamic W', 'saif_dynamic'), (25, 'SAIF total W', 'saif_total')):
                cur = ws.cell(r, col).value
                prop = num(s[key])
                if cur is None or abs(float(cur) - prop) > 1e-9:
                    changes.append(('T4', ws.title, f'{get_column_letter(col)}{r}', f'{lab} ({block})', fld, cur, prop,
                                    ('invalid capture (idle core) replaced by PC-verified re-capture'
                                     if (v, block) == ('RV64IM_8SP', 'coremark') else
                                     s['source'] + '; SAIF tables reconciled on the final implementation'),
                                    'T4/T4_saif_power.md'))
    # 3. SAIF detail sheet: total and dynamic per row
    ws = mwb['SAIF detail']
    block = None
    for r in range(1, ws.max_row + 1):
        lab = ws.cell(r, 2).value
        if lab in ('dhrystone', 'coremark'):
            block = lab
            continue
        if not block or not lab or not str(lab).startswith('RV'):
            continue
        s = saifv.get((lab, block))
        if not s:
            continue
        for col, fld, key in ((3, 'total W', 'saif_total'), (4, 'dynamic', 'saif_dynamic')):
            cur = ws.cell(r, col).value
            prop = num(s[key])
            if cur is None or abs(float(cur) - prop) > 1e-9:
                changes.append(('T4', 'SAIF detail', f'{get_column_letter(col)}{r}', f'{lab} ({block})', fld, cur, prop,
                                'replace the whole row with T4 authoritative (all components, coverage, validity)',
                                'T4/T4_saif_authoritative.csv'))
    # 4. Normalized metrics
    ws = mwb['Normalized metrics']
    for r in range(6, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if not v or v not in nm:
            continue
        n = nm[v]
        prop = {4: num(n['true_coremark_its']), 5: round(float(cmv[v]['TRUE CoreMark/MHz']), 4),
                10: num(n['cm_total_W']), 11: num(n['CM_per_W']), 12: num(n['mJ_per_CM_iter']),
                13: num(n['true_coremark_its'])}
        names = {4: 'CoreMark', 5: 'CM/MHz', 10: 'total W', 11: 'CM/W', 12: 'mJ/iter', 13: 'iters/s'}
        for col, p in prop.items():
            cur = ws.cell(r, col).value
            if cur is None or abs(float(cur) - float(p)) > 1e-6 * max(abs(float(p)), 1.0):
                changes.append(('T3/T4', 'Normalized metrics', f'{get_column_letter(col)}{r}', v, names[col], cur, p,
                                'true CoreMark it/s (T3) and authoritative SAIF power (T4)',
                                'T4/T4_normalized_metrics.csv'))
        lut_core = num(n.get('core_LUT'))
        if lut_core:
            cur = ws.cell(r, 7).value
            p = round(float(n['true_coremark_its']) / lut_core, 5)
            if cur is None or abs(float(cur) - p) > 1e-6 * max(p, 1.0):
                changes.append(('T3', 'Normalized metrics', f'G{r}', v, 'CM/LUT core', cur, p,
                                'recomputed from true CoreMark it/s', 'T3/T3_fragment_corrected_coremark.csv'))
    # 5. Embench simulation: huffbench / slre / wikisort
    ws = mwb['Embench simulation']
    colof = {ws.cell(5, c).value: c for c in range(3, ws.max_column + 1)}
    rowof = {ws.cell(r, 2).value: r for r in range(6, ws.max_row + 1)}
    for e in e12:
        c, r = colof.get(e['benchmark']), rowof.get(e['variant'])
        if c and r:
            cur = ws.cell(r, c).value
            prop = num(e['cycles'])
            if cur != prop:
                changes.append(('T5', 'Embench simulation', f'{get_column_letter(c)}{r}', e['variant'], e['benchmark'],
                                cur, prop, '4 KB stack overflow in the port; re-measured with a 12 KB stack (verify OK)',
                                'T5/T5_embench_no_result.md'))
    # 6. Clocking resources: BUFGCTRL / PLL from the utilisation reports
    ws = mwb['Clocking resources']
    for r in range(6, ws.max_row + 1):
        v, b = ws.cell(r, 2).value, ws.cell(r, 3).value
        if not v or not b:
            continue
        u = f'{ENV}/logs/final_impl/{v}_{b}/utilization.rpt'
        if not os.path.exists(u):
            continue
        t = open(u).read()
        for col, lab in ((4, 'BUFGCTRL'), (5, 'MMCME2_ADV'), (6, 'PLLE2_ADV')):
            m = re.search(rf'^\|\s*{lab}\s*\|\s*(\d+)', t, re.M)
            rep = int(m.group(1)) if m else 0
            cur = ws.cell(r, col).value
            if cur != rep:
                changes.append(('T7', 'Clocking resources', f'{get_column_letter(col)}{r}', f'{v} ({b})', lab, cur, rep,
                                'value in logs/final_impl/<build>/utilization.rpt', 'T7/T7_metadata_checks.md'))
    # 7. Text items
    txt = [
        ('T5', 'Verification', 'note / Embench column', 'all', 'huffbench/slre/wikisort wording',
         '"400M-cycle simulation cap, not a functional failure"',
         '4 KB stack overflow in the Embench port; with a 12 KB stack all three complete (verify OK) on all 16 variants',
         'stack audit + causal relink', 'T5/T5_embench_no_result.md'),
        ('T5', 'Embench simulation', 'huffbench column (RV32 rows)', 'RV32 variants', 'huffbench',
         'measured with the stack overwriting its heap', 'replace with the 12 KB-stack cycles (listed above)',
         'stack entered .bss by 3,744 B; instret changes by -8.6%', 'T5/T5_embench_no_result.md'),
        ('T5', 'Embench simulation', 'xgboost column', 'all', 'xgboost', '"ok"',
         'annotate: verify_benchmark is vacuous at GLOBAL_SCALE_FACTOR=1; code exceeds the 32 KB IMEM',
         'upstream Embench comma-operator bug', 'T5/T5_embench_no_result.md'),
        ('T6', 'Verification / Contents', 'randomized-testing note', 'all 7/8-stage', 'AAPG divergence',
         '"control-flow divergence ... affects no reported measurement" (asserted)',
         'root cause: IF_IO_Register skips the predicted-taken squash when an M-extension op is in ID; '
         'TRIG = 0 on every reported workload incl. the full board binaries (proven)',
         'testbench trigger monitor, 2,100 + 20 runs', 'T6/T6_aapg_divergence_scope.md'),
        ('T6', 'Verification', 'note', 'RV64IM_7SP', 'secondary defects',
         '"passes every suite"', 'add: wrong-path commit once per CoreMark iteration; wrong byte loads in qrduino (sim only)',
         'P1/P3 and lock-step monitors', 'T6/T6_aapg_divergence_scope.md'),
        ('T8', 'CPI and stalls', 'load-use hazard % (4 x 8SP rows)', '8-stage variants', 'load-use hazard',
         'one column (~18% on 8SP)', 'two columns: load-use ~1.9% and execution-use ~8.6% (RV64 Dhrystone)',
         'Hazard_Unit.v:124-125 ORs them into the profiled signal', 'T8/T8_stall_breakdown.md'),
        ('T8', 'CPI and stalls', 'stall columns', 'all', 'stall percentages',
         'overlapping occupancies (pc/front-end/ID-EX identical, not additive)',
         'add the new additive sheet "Stall breakdown by cause"', 'partition verified exact on 112/112 runs',
         'T8/t8_stall_breakdown.csv'),
        ('T9', 'Microkernel penalties', 'flush penalty per depth', 'all', 'mispredict penalty',
         'one value per depth (2/3/4/5)', 'two directions: NT->T 2/3/3/4 and T->NT 2/3/5/6, plus a 2-cycle '
         'correctly-predicted-taken cost from 7 stages on', 'matched pairs + exact least-squares fit (residual 0.0%)',
         'T9/T9_mispredict_penalty.md'),
        ('T12', 'Verification', 'randomized (AAPG) column', 'all', 'AAPG pass counts',
         '7/20, 3/20, 13 mismatch, ...', 'every completing run is byte-identical to the 5-stage core on the same ROM image '
         '(113/113, 0 differing words); the rest are non-terminations from the single T6 defect',
         'the MISMATCHes were Sail\'s 0x80000000 RAM base vs the DUT\'s 0x0', 'T12/T12_aapg_address_matched.md'),
        ('T12', 'Contents', 'known-gaps note', '-', 'randomized failure modes',
         'two failure modes (timeout and mismatch)', 'one: non-termination; the mismatch mode is a reference-map artefact',
         'DUT-vs-DUT comparison', 'T12/T12_aapg_address_matched.md'),
        ('T7', 'Verification', 'riscv-tests note', 'all', 'ma_data',
         '"these cores trap"', 'loads/stores trap, except RV64 ld/lwu (silently aligned) and a misaligned sh on '
         'RV32IM_5SP/8SP/8SP_withoutOpt (traps but still writes)', 'direct probe on all 16', 'T7/T7_metadata_checks.md'),
        ('T7', 'RV32/RV64 SoC and FPGA', 'notes column', 'all SoC rows', 'clock primitive',
         '"the MMCM\'s exact output"', '"the PLL\'s exact output" (PLLE2_ADV in all 32 builds)',
         'utilization.rpt', 'T7/T7_metadata_checks.md'),
        ('T7', 'Embench FPGA RV64', 'IM_5SP nettle-aes build MHz', 'IM_5SP', 'nettle-aes',
         '30 (unexplained)', 'keep 30, add footnote: did not close at 38/37 MHz (WNS -0.131/-0.106); rebuilt at 30',
         'embench_retry2.log, MANIFEST.csv', 'T7/T7_metadata_checks.md'),
        ('T1', 'Embench FPGA RV32', 'all 40 rows', 'all RV32', 'FPGA cycles/instret',
         'awaiting re-measurement (RV32 counter fix)', 'measured on the 40 rebuilt bitstreams — see sheet T1 Sim ref '
         'and T1_fragment_embench_fpga_rv32.csv; 40/40 EXACT against RTL simulation in both counters',
         'board runs 09-23 on the 09-22 (v4 counter read) bitstreams', 'T1/T1_rv32_embench_fpga.md'),
        ('T1', 'Embench FPGA vs simulation', 'B44', 'summary line', 'row count',
         '40 of 40 rows match simulation exactly in BOTH counters',
         '80 of 80 rows match simulation exactly in BOTH counters (40 RV64 + 40 RV32)',
         'the 40 RV32 cross-check rows are now measured', 'T1/T1_fragment_fpga_vs_sim_rows.csv'),
        ('T1', 'Embench FPGA vs simulation', 'B45', 'exclusion note', 'note',
         'RV32 is excluded: those runs predate the RV32 counter-read fix and their cycle counts are invalid',
         'RV32 rows are the rebuilt (v4 counter-read) bitstreams, run 2026-09-23; both widths are complete',
         'the pre-fix runs are superseded, not used', 'T1/T1_rv32_embench_fpga.md'),
        ('T1', 'Embench FPGA vs simulation', 'rows 43+', '40 RV32 rows', 'new rows',
         'absent', 'append T1_fragment_fpga_vs_sim_rows.csv (40 rows, same column order)',
         '40/40 EXACT in both counters', 'T1/T1_fragment_fpga_vs_sim_rows.csv'),
    ]
    changes += txt

    ws = wb.create_sheet('Change list')
    title(ws, 'Change list — every master-file cell a fragment changes',
          f'Master: {os.path.basename(MASTER)} (not modified).  One row per cell; text items at the end.  '
          'Yellow = proposed value.  Review, then apply.')
    hdr = ['#', 'task', 'master sheet', 'cell', 'row', 'field', 'current value', 'proposed value', 'reason', 'evidence']
    rows = [[i] + list(c) for i, c in enumerate(changes, 1)]
    first, last = table(ws, 4, hdr, rows, widths={'reason': 60, 'evidence': 34, 'current value': 22, 'proposed value': 30,
                                                     'row': 24, 'master sheet': 22})
    for r in range(first, last + 1):
        ws.cell(r, 8).fill = CHG_FILL
        for c in (7, 8, 9):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical='top')

    # ---------------------------------------------------------------- T3 CoreMark (formulas)
    ws = wb.create_sheet('T3 CoreMark')
    title(ws, 'T3 — CoreMark and Dhrystone per iteration, true scores',
          'Inputs in blue: RTL-simulation cycles/instructions per iteration (difference of 10 vs 20 iterations), run clock, '
          'master value.  Black cells are formulas.  Source: T3/T3_fragment_corrected_coremark.csv')
    hdr = ['variant', 'run MHz', 'image MHz', 'cycles/iter', 'instr/iter', 'CPI', 'board iterations', 'board int seconds',
           'replayed print (it/s)', 'master it/s', 'replay = master', 'TRUE it/s', 'TRUE CM/MHz', 'master error',
           'DH cycles/iter', 'DH instr/iter', 'DH run MHz', 'predicted Dhrystones/s', 'master Dhrystones/s', 'DH diff']
    rows = []
    for i, r in enumerate(cm, 5):
        rows.append([r['variant'], num(r['CoreMark run MHz']), num(r['image compiled for MHz']),
                     num(r['CoreMark cycles/iter (sim)']), num(r['CoreMark instr/iter']), f'=D{i}/E{i}',
                     num(r['board iterations (replayed)']), num(r['board int seconds (replayed)']), f'=INT(G{i}/H{i})',
                     num(r['master it/s']), f'=IF(I{i}=J{i},"yes","NO")', f'=B{i}*1000000/D{i}', f'=L{i}/B{i}',
                     f'=(J{i}-L{i})/L{i}', num(r['Dhrystone cycles/iter']), num(r['Dhrystone instr/iter']),
                     num(r['Dhrystone run MHz']), f'=Q{i}*1000000/O{i}', num(r['master Dhrystones/s']), f'=(S{i}-R{i})/R{i}'])
    first, last = table(ws, 4, hdr, rows, number_formats={'CPI': '0.0000', 'TRUE it/s': '0.00', 'TRUE CM/MHz': '0.0000',
                                                          'master error': '+0.00%;-0.00%', 'predicted Dhrystones/s': '#,##0.0',
                                                          'DH diff': '+0.000%;-0.000%', 'cycles/iter': '#,##0.0',
                                                          'instr/iter': '#,##0.0'})
    for r in range(first, last + 1):
        for c in (2, 3, 4, 5, 7, 8, 10, 15, 16, 17, 19):
            ws.cell(r, c).font = Font(name=F, size=9, color='FF0000FF')
        if ws.cell(r, 1).value == 'RV32IM_8SP_withoutOpt':
            for c in (1, 3, 10, 14):
                ws.cell(r, c).fill = BAD_FILL
    n = last + 2
    ws.cell(n, 1, 'Notes').font = Font(name=F, bold=True, size=9)
    for k, t in enumerate([
        'board int seconds = floor(ticks / image clock); the ports use HAS_FLOAT 0, so CoreMark prints iterations / int_seconds.',
        'Replaying core_main.c calibration with the simulated cycles/iteration reproduces the printed board value on 16/16 (column K).',
        'RV32IM_8SP_withoutOpt: the "114 MHz" image is byte-identical to the 100 MHz build, so the image clock is 100 MHz.',
        'Report CoreMark as TRUE it/s = f x iterations / cycles (or N*f/Total_ticks from the UART), not the integer Iterations/Sec line.']):
        ws.cell(n + 1 + k, 1, t).font = Font(name=F, size=9)

    # ---------------------------------------------------------------- T4 normalized metrics (formulas)
    ws = wb.create_sheet('T4 Norm metrics')
    title(ws, 'T4 — Normalised metrics with true CoreMark and authoritative SAIF power',
          'Blue = inputs (true CoreMark from T3, SAIF power from T4 authoritative, Dhrystone/s and core LUT from the master). '
          'Cross-variant energy should use the Dhrystone columns: CoreMark windows sample different phases per variant.')
    hdr = ['variant', 'true CoreMark it/s', 'Dhrystones/s', 'core LUT', 'CM SAIF total W', 'CM core-attr W',
           'DH SAIF total W', 'DH core-attr W', 'CM/W', 'mJ per CM iter', 'CM/W core-attr', 'mJ per CM iter core-attr',
           'uJ per Dhry iter', 'uJ per Dhry iter core-attr', 'CM per kLUT core']
    rows = []
    for i, v in enumerate(sorted(nm), 5):
        cmr, dhr = saifv[(v, 'coremark')], saifv[(v, 'dhrystone')]
        rows.append([v, num(nm[v]['true_coremark_its']), num(nm[v]['dhrystones_per_s']), num(nm[v]['core_LUT']),
                     num(cmr['saif_total']), num(cmr['core_attributable_dyn']), num(dhr['saif_total']),
                     num(dhr['core_attributable_dyn']), f'=B{i}/E{i}', f'=1000*E{i}/B{i}', f'=B{i}/F{i}',
                     f'=1000*F{i}/B{i}', f'=1000000*G{i}/C{i}', f'=1000000*H{i}/C{i}', f'=1000*B{i}/D{i}'])
    first, last = table(ws, 4, hdr, rows, number_formats={h: '0.000' for h in hdr[4:8]} | {
        'CM/W': '0.0', 'mJ per CM iter': '0.000', 'CM/W core-attr': '0.0', 'mJ per CM iter core-attr': '0.000',
        'uJ per Dhry iter': '0.000', 'uJ per Dhry iter core-attr': '0.000', 'CM per kLUT core': '0.00',
        'true CoreMark it/s': '0.00', 'Dhrystones/s': '#,##0'})
    for r in range(first, last + 1):
        for c in range(2, 9):
            ws.cell(r, c).font = Font(name=F, size=9, color='FF0000FF')

    # ---------------------------------------------------------------- CSV fragments
    csv_sheet(wb, 'T4 SAIF', f'{REV}/T4/T4_saif_authoritative.csv',
              'Authoritative SAIF power: every capture applied to the FINAL implementation; coverage, validity, window.')
    csv_sheet(wb, 'T4 Windows', f'{REV}/T4/capture_window_map.csv',
              'Capture window (1200-1300 us) mapped to core cycles, instret and functions (RTL simulation, cycle-exact).')
    csv_sheet(wb, 'T4 Validity', f'{REV}/T4/capture_validity_dmem_activity.csv',
              'Toggle rate per us inside the CPU and its memories for the 32 original captures; RV64IM_8SP_coremark is idle '
              '(replaced by the v2 re-capture).')
    csv_sheet(wb, 'T3 SoC reports', f'{REV}/T3/soc_rows_from_vivado_reports.csv',
              'Every SoC row re-read from its Vivado reports (report date and md5 for provenance).')
    csv_sheet(wb, 'T3 Image audit', f'{REV}/T3/image_audit_32builds.tsv',
              'Image each SoC project loads vs a fresh build at its PLL clock: 31/32 byte-identical.', header=False,
              names=['project', 'image', 'PLL MHz', 'build Hz', 'project md5', 'rebuilt md5', 'verdict'])
    csv_sheet(wb, 'T5 Embench 12KB', f'{REV}/T5/T5_fragment_embench_12KBstack.csv',
              'huffbench/slre/wikisort rebuilt with a 12 KB stack: 48/48 verify OK.')
    csv_sheet(wb, 'T5 Stack audit', f'{REV}/T5/stack_audit_19x4.tsv',
              'Lowest sp reached vs _bss_end and the RAM base, 19 benchmarks x 4 ISAs.', header=False,
              names=['variant', 'isa', 'benchmark', 'halted', 'bss_end', 'stack_end', 'min_sp', 'depth_B',
                     'bytes_into_bss', 'verdict'])

    # T6 summary (pivot) + raw
    trig = rcsv(f'{REV}/T6/trigger_counts_all_workloads.csv', False, ['variant', 'workload', 't6_trig', 'cycles', 'end'])
    board = rcsv(f'{REV}/T6/trigger_counts_board_binaries.csv', False, ['variant', 'workload', 't6_trig', 'cycles', 'end'])
    ws = wb.create_sheet('T6 Trigger summary')
    title(ws, 'T6 — TRIG (predicted-taken squash suppressed by an M-op in ID) per variant and workload class',
          'Cells: runs with TRIG > 0 / runs.  TRIG = 0 on every reported workload; fires on all 127 AAPG timeouts.')
    cls = lambda w: ('AAPG' if w.startswith('aapg_') else 'micro-kernels' if w.startswith('mk_') else
                     'Embench (19)' if w.startswith('embench') else 'CoreMark 10 it' if w.startswith('coremark') else
                     'Dhrystone 1000 it')
    C = ['CoreMark board binary', 'Dhrystone board binary', 'CoreMark 10 it', 'Dhrystone 1000 it', 'Embench (19)',
         'micro-kernels', 'AAPG']
    agg = {}
    for r in trig:
        k = (r['variant'], cls(r['workload']))
        a = agg.setdefault(k, [0, 0]); a[0] += int(r['t6_trig']) > 0; a[1] += 1
    for r in board:
        k = (r['variant'], 'CoreMark board binary' if 'Coremark' in r['workload'] else 'Dhrystone board binary')
        a = agg.setdefault(k, [0, 0]); a[0] += int(r['t6_trig']) > 0; a[1] += 1
    V = sorted({r['variant'] for r in trig})
    table(ws, 4, ['variant'] + C, [[v] + [f'{agg.get((v, c), [0, 0])[0]} / {agg.get((v, c), [0, 0])[1]}' for c in C] for v in V],
          widths={c: 20 for c in C})
    csv_sheet(wb, 'T6 TRIG all runs', f'{REV}/T6/trigger_counts_all_workloads.csv',
              '2,100 runs (10 deep variants).', header=False, names=['variant', 'workload', 't6_trig', 'cycles', 'end'])
    csv_sheet(wb, 'T6 TRIG board', f'{REV}/T6/trigger_counts_board_binaries.csv',
              'Full board binaries run to completion (cap = expected length + 5%).', header=False,
              names=['variant', 'workload', 't6_trig', 'cycles', 'end'])
    csv_sheet(wb, 'T6 Lockstep', f'{REV}/T6/p2_lockstep.csv',
              'Value-level lock-step vs the 5SP (P2).  RV32 deep-variant workload rows superseded by the next sheet.')
    csv_sheet(wb, 'T6 Lockstep RV32 w64', f'{REV}/T6/p2_lockstep_rv32_win64.csv',
              'RV32 7SP_BRAM/_Opt/8SP/withoutOpt re-run with a 64-entry resync window: 0 value divergences, all complete.')
    csv_sheet(wb, 'T6 P1 workloads', f'{REV}/T6/p1_controlflow_workloads.csv',
              'Committed-stream continuity violations (P1) per workload.')

    # T7 table
    ws = wb.create_sheet('T7 Metadata')
    title(ws, 'T7 — Metadata consistency checks', 'Details and evidence: T7/T7_metadata_checks.md')
    t7 = [
        (1, 'RV64 IM_5SP nettle-aes at 30 MHz', 'confirmed', '38 MHz WNS -0.131, 37 MHz WNS -0.106; rebuilt at 30 MHz WNS +0.735. Runtime 175.532 ms correct for that bitstream; footnote.'),
        (2, 'Core-only tops / identical 64.637 MHz', 'confirmed genuine', 'Top name RV64IM72F6SP_CORE inherited; each run synthesises its own sources (7SP runs contain IF_IO). Incremental synthesis not used. Same EX critical path in 6SP/7SP/7SP_BRAM.'),
        (3, 'RV32 IM_8SP static 0.136 W', 'estimator artefact', 'Identical environment, Tj, process, operating conditions, pins/standards/banks; the 17 mW is Vcco33 static only, reproducible on re-run.'),
        (4, 'PLL vs MMCM', 'corrected', 'PLLE2_ADV = 1, MMCME2_ADV = 0 in all 32; RV64 IM_8SP PLL cell 0 -> 1; notes say MMCM -> PLL.'),
        (5, 'Vivado version / strategies', 'confirmed', 'v2025.2 (6299465); Flow_PerfOptimized_high + Performance_ExplorePostRoutePhysOpt; INCREMENTAL_DISABLED in all 32 logs; built 09-15 18:06 to 09-16 14:26.'),
        (6, 'RISCOF reference', 'confirmed', 'sail_cSim (Sail 0.6, live); riscv-arch-test old-framework-2.x @ 6f7f47bd (2.7.4) with one local arch_test.h edit (disclose).'),
        (7, 'riscv-tests ma_data', 'confirmed + refined', 'Only ma_data fails. Loads/stores trap (cause 4/6) except RV64 ld/lwu (silently aligned) and misaligned sh on RV32IM_5SP/8SP/8SP_withoutOpt (traps but writes). No reported workload performs a misaligned access.'),
    ]
    first, last = table(ws, 4, ['#', 'item', 'verdict', 'finding'], [list(x) for x in t7], widths={'finding': 110, 'item': 38})
    for r in range(first, last + 1):
        ws.cell(r, 4).alignment = Alignment(wrap_text=True, vertical='top')

    csv_sheet(wb, 'T8 Stall by cause', f'{REV}/T8/t8_stall_breakdown.csv',
              'Every cycle attributed to exactly one cause (112 runs; partition exact and non-perturbing on all).')
    csv_sheet(wb, 'T9 Mispredict penalty', f'{REV}/T9/t9_mispredict_penalties.csv',
              'Matched-pair and least-squares penalties by direction (48 families, residual 0.0%).')
    csv_sheet(wb, 'T12 AAPG vs 5SP', f'{REV}/T12/t12_aapg_dutvsdut.csv',
              'AAPG signatures compared with the 5-stage core on the identical ROM image (no link-address artefact).')
    csv_sheet(wb, 'T10 Dynamic mix', f'{REV}/T10/t10_dynamic_mix.csv',
              'Instructions per iteration and dynamic mix (difference of two iteration counts, RTL simulation).')
    csv_sheet(wb, 'T13 Worst paths', f'{REV}/T13_worst_paths_soc.tsv',
              'Worst setup path of each SoC build (logs/final_impl/*/timing.rpt).')

    # T1 simulation reference with runtime formula placeholders
    ws = wb.create_sheet('T1 Sim ref')
    title(ws, 'T1 — RV32 Embench: the 40 rebuilt bitstreams, board vs RTL simulation',
          'FPGA columns are the measured board values (09-23); the EXACT checks are formulas over both counters.')
    ref = rcsv(f'{REV}/T1/embench_rv32_simref_v4.csv')
    board = {(r['variant'], r['benchmark']): r for r in rcsv(f'{REV}/T1/T1_fragment_fpga_vs_sim_rows.csv')}
    clk = {'RV32I_5SP': 45, 'RV32IM_5SP': 43, 'RV32IM_6SP': 50, 'RV32IM_7SP': 57.501, 'RV32IM_7SP_BRAM': 71.999,
           'RV32IM_7SP_BRAM_Opt': 92, 'RV32IM_8SP_withoutOpt': 113.999, 'RV32IM_8SP': 122}
    hdr = ['variant', 'benchmark', 'sim cycles', 'sim instret', 'sim CPI', 'sim verify', 'build MHz', 'FPGA cycles',
           'FPGA instret', 'cycles match', 'instret match', 'runtime ms (FPGA)', 'image md5', 'ELF sha256', 'bitstream']
    rows = []
    for i, r in enumerate(ref, 5):
        b = board.get((r['variant'], r['benchmark']), {})
        rows.append([r['variant'], r['benchmark'], num(r['sim_cycles']), num(r['sim_instret']), f'=C{i}/D{i}',
                     r['sim_verify'], clk[r['variant']], num(b.get('FPGA cycles')), num(b.get('FPGA instret')),
                     f'=IF(H{i}="","",IF(H{i}=C{i},"EXACT","MISMATCH"))', f'=IF(I{i}="","",IF(I{i}=D{i},"EXACT","MISMATCH"))',
                     f'=IF(H{i}="","",H{i}/G{i}/1000)', r['image_md5'], r['elf_sha256'], r['bitstream']])
    first, last = table(ws, 4, hdr, rows, widths={'ELF sha256': 20, 'bitstream': 50, 'image md5': 20},
                        number_formats={'sim CPI': '0.0000', 'sim cycles': '#,##0', 'sim instret': '#,##0',
                                        'runtime ms (FPGA)': '0.000'})
    # ---------------------------------------------------------------- Index
    title(idx, 'RV-IM100 revision — review workbook (P0 fragments)',
          f'Built {datetime.date.today().isoformat()} from revision_0922/.  The master file is not modified.')
    items = [
        ('Change list', 'every master cell a fragment changes: current vs proposed, reason, evidence', 'all', f'{len(changes)} items'),
        ('T3 CoreMark', 'true CoreMark / Dhrystone per iteration; replay of the printed board values (formulas)', 'RV32/RV64 SoC and FPGA, Normalized metrics', 'done'),
        ('T3 SoC reports', 'all 32 SoC rows from the Vivado reports', 'RV32/RV64 SoC and FPGA', 'done: only PLL cell differs'),
        ('T3 Image audit', 'image provenance of the 32 SoC projects', '-', 'done: 31/32 match'),
        ('T4 Norm metrics', 'normalised metrics recomputed (formulas), incl. uJ/Dhrystone and core-attributable', 'Normalized metrics', 'done'),
        ('T4 SAIF', 'authoritative SAIF table (final implementations), coverage, validity', 'SAIF detail + SoC SAIF columns', 'done: 32/32 valid'),
        ('T4 Windows', 'capture window -> cycles, instret, functions', '-', 'done'),
        ('T4 Validity', 'toggle activity per capture (idle-core check)', '-', 'done'),
        ('T5 Embench 12KB', 'huffbench/slre/wikisort with a 12 KB stack', 'Embench simulation', 'done: 48/48 OK'),
        ('T5 Stack audit', 'stack depth vs allocation', 'Verification wording', 'done'),
        ('T6 Trigger summary', 'AAPG defect trigger count per variant/workload', 'Verification wording', 'done: 0 on all reported'),
        ('T6 TRIG all runs / board', 'raw trigger counts', '-', 'done'),
        ('T6 Lockstep (+RV32 w64)', 'value-level comparison vs 5SP', '-', 'done'),
        ('T6 P1 workloads', 'control-flow continuity per workload', '-', 'done'),
        ('T7 Metadata', 'seven metadata checks', 'Clocking, SoC notes, Verification', 'done'),
        ('T10 Dynamic mix', 'instructions per iteration and mix', 'new table', 'done'),
        ('T13 Worst paths', 'worst setup path per SoC build', 'Table II / text', 'done'),
        ('T1 Sim ref', 'RV32 Embench: measured board cycles/instret vs the RTL simulation reference', 'Embench FPGA RV32', 'done: 40/40 EXACT'),
        ('T8 Stall by cause', 'every cycle attributed to one cause; load-use vs execution-use separated', 'CPI and stalls (+ new sheet)', 'done: 112/112 exact'),
        ('T9 Mispredict penalty', 'misprediction penalty by direction, matched pairs', 'Microkernel penalties', 'done: residual 0.0%'),
        ('T12 AAPG vs 5SP', 'randomized programs vs the reference core, no link-address artefact', 'Verification', 'done: 113/113 identical'),
        ('(paper_numbers.csv)', 'every number the paper quotes, with its source; figures/ holds the 13 regenerated PDFs', 'text, tables, figures', 'done: 1,561 numbers'),
    ]
    table(idx, 4, ['sheet', 'contents', 'master sheet affected', 'status'], [list(x) for x in items],
          widths={'contents': 80, 'master sheet affected': 40, 'status': 28, 'sheet': 26})
    n = 5 + len(items) + 1
    for k, t in enumerate(['Legend', 'Yellow fill: proposed value (Change list).',
                           'Orange fill: invalid or contradicted master value.', 'Blue text: input value; black: formula.',
                           'Written reports with full reasoning: revision_0922/T*/*.md']):
        c = idx.cell(n + k, 1, t)
        c.font = Font(name=F, bold=(k == 0), size=9)
    idx.freeze_panes = None

    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 90
    wb.save(OUT)
    print('wrote', OUT, 'sheets:', len(wb.worksheets), 'change items:', len(changes))


if __name__ == '__main__':
    main()
