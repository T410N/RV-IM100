#!/usr/bin/env python3
"""T8: run the cause-based stall accounting and check its invariants.

Dhrystone and CoreMark use each variant's own board image and the same 10M-retired-
instruction window as the master `CPI and stalls` sheet; the five FPGA Embench
benchmarks run to completion.

Two checks per run:
  * exact partition:  sum(buckets) + retired-cycles == cycles
  * no perturbation:  cycles and retired equal the production simulator's

and, for Dhrystone/CoreMark, cycles/retired are compared with the master sheet.

Writes logs/t8_stall_breakdown.csv.
"""
import csv, glob, json, os, re, subprocess, sys
from concurrent.futures import ThreadPoolExecutor

ENV = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
ROOT = os.path.dirname(ENV)
TMP = os.environ.get('CFTMP', '/tmp/rv-im100')
VARIANTS = ['RV32I_5SP', 'RV32IM_5SP', 'RV32IM_6SP', 'RV32IM_7SP', 'RV32IM_7SP_BRAM', 'RV32IM_7SP_BRAM_Opt',
            'RV32IM_8SP_withoutOpt', 'RV32IM_8SP', 'RV64I_5SP', 'RV64IM_5SP', 'RV64IM_6SP', 'RV64IM_7SP',
            'RV64IM_7SP_BRAM', 'RV64IM_7SP_BRAM_Opt', 'RV64IM_8SP_withoutOpt', 'RV64IM_8SP']
EMBENCH = ['matmult-int', 'crc32', 'nettle-aes', 'statemate', 'md5sum']
INSTR = 10_000_000


def prof(exe, args, tag):
    p = f'{TMP}/t8_{tag}.prof'
    subprocess.run([exe] + args + [f'+PROF={p}'], capture_output=True)
    d = {}
    for line in open(p):
        w = line.split()
        if len(w) >= 2:
            d[w[0]] = int(w[1])
        if len(w) == 4:
            d[w[2]] = int(w[3])
    os.remove(p)
    return d


def board_image(v, bench):
    proj = f'{ROOT}/RV-IM100_RTL/project_files/RV{v[2:4]}s/SoCs/{v}_{"Dhry" if bench == "dhrystone" else "Coremark"}'
    xpr = glob.glob(f'{proj}/*.xpr')[0]
    base = os.path.splitext(xpr)[0]
    files = [f.replace('$PPRDIR', proj).replace('$PSRCDIR', base + '.srcs')
             for f in re.findall(r'<File Path="([^"]+)"', open(xpr).read())]
    im = [f for f in files if f.endswith('Instruction_Memory.v')][0]
    name = os.path.basename(re.findall(r'readmemh\(\s*"([^"]+)"', re.sub(r'//.*', '', open(im).read()))[0])
    return [f for f in files if os.path.basename(f) == name and os.path.exists(f)][0]


def one(v, bench):
    t8 = f'{ENV}/build_t8/socs_{v}/Vsim_top'
    ref = f'{ENV}/build/socs_{v}/Vsim_top'
    if bench in ('dhrystone', 'coremark'):
        img = board_image(v, bench)
        args = [f'+HEX={img}', '+NOHALT', f'+MAX_INSTR={INSTR}', '+MAX_CYCLES=200000000']
    else:
        isa = f'rv{v[2:4]}{"im" if "IM" in v else "i"}'
        args = [f'+HEX={ENV}/build/embench_fpga/{isa}/{bench}.mem', '+HALT_ADDR=10012000', '+MAX_CYCLES=400000000']
    a = prof(t8, args, f'{v}_{bench}')
    b = prof(ref, args, f'{v}_{bench}_ref')
    buckets = {k[3:]: n for k, n in a.items() if k.startswith('t8_')}
    total = sum(buckets.values())
    row = dict(variant=v, benchmark=bench, cycles=a['cycles'], retired=a['retired'],
               cpi=round(a['cycles'] / a['retired'], 4) if a['retired'] else '',
               partition_exact=(total == a['cycles']),
               same_as_production_sim=(a['cycles'] == b['cycles'] and a['retired'] == b['retired']))
    for k, n in buckets.items():
        row[k] = n
        if k != 'retired_cycles':
            row[k + '_pct'] = round(100 * n / a['cycles'], 3)
    return row


def main():
    jobs = [(v, b) for v in VARIANTS for b in ['dhrystone', 'coremark'] + EMBENCH]
    with ThreadPoolExecutor(int(os.environ.get('JOBS', '6'))) as ex:
        rows = list(ex.map(lambda t: one(*t), jobs))
    keys = ['variant', 'benchmark', 'cycles', 'retired', 'cpi', 'partition_exact', 'same_as_production_sim']
    order = ['retired_cycles', 'div_busy', 'mul_busy', 'load_use', 'exec_use', 'csr_not_ready', 'mispredict_flush',
             'taken_refill', 'jump_redirect', 'stall_other', 'frontend_empty']
    keys += [k for n in order for k in (n, n + '_pct') if any(k in r for r in rows)]
    out = f'{ENV}/logs/t8_stall_breakdown.csv'
    with open(out, 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=keys, extrasaction='ignore')
        w.writeheader()
        w.writerows(rows)
    bad = [r for r in rows if not r['partition_exact'] or not r['same_as_production_sim']]
    print(f'{len(rows)} runs -> {out}')
    print(f'partition exact: {sum(r["partition_exact"] for r in rows)}/{len(rows)}; '
          f'identical to the production simulator: {sum(r["same_as_production_sim"] for r in rows)}/{len(rows)}')
    for r in bad:
        print('  !!', r['variant'], r['benchmark'], r['partition_exact'], r['same_as_production_sim'])
    # cross-check against the master CPI sheet
    try:
        import openpyxl
        wb = openpyxl.load_workbook(f'{ROOT}/Research_data_MASTER_0921.xlsx', data_only=True)
        ws = wb['CPI and stalls']
        block = None
        n_ok = n_tot = 0
        for r in range(1, ws.max_row + 1):
            h = ws.cell(r, 2).value
            if h in ('dhrystone', 'coremark'):
                block = h
                continue
            if not block or not h or not str(h).startswith('RV'):
                continue
            mine = next((x for x in rows if x['variant'] == h and x['benchmark'] == block), None)
            if mine and ws.cell(r, 3).value:
                n_tot += 1
                same = int(float(ws.cell(r, 3).value)) == mine['cycles'] and int(float(ws.cell(r, 4).value)) == mine['retired']
                n_ok += same
                if not same:
                    print(f"  sheet mismatch {h} {block}: sheet {ws.cell(r,3).value}/{ws.cell(r,4).value} "
                          f"vs {mine['cycles']}/{mine['retired']}")
        print(f'cycles+retired match the master CPI sheet: {n_ok}/{n_tot}')
    except Exception as e:
        print('sheet cross-check skipped:', e)


if __name__ == '__main__':
    main()
