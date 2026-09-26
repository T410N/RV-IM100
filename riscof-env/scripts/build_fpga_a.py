#!/usr/bin/env python3
"""~/Documents/Research_data_FPGA-A.xlsx -- FPGA-measured results only.

One row per bitstream, because each (variant, benchmark) pair has its own clock.
Earlier measurements are carried over ONLY where they were taken at the clock the
current bitstream actually uses; where the clock has since changed the cell is left
blank and highlighted, because that number no longer describes the hardware.
"""
import csv, pathlib, re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ENV = pathlib.Path(__file__).resolve().parent.parent
DST = pathlib.Path.home() / "Documents" / "Research_data_FPGA-A.xlsx"
man = {(r["variant"], r["benchmark"]): r
       for r in csv.DictReader((ENV.parent / "bitstream" / "MANIFEST.csv").open())}
fb = {r["variant"]: r for r in csv.DictReader((ENV / "logs" / "fpga_benchmarks.csv").open())}

ORDER = {"RV32": ["RV32I_5SP", "RV32IM_5SP", "RV32IM_6SP", "RV32IM_7SP", "RV32IM_7SP_BRAM",
                  "RV32IM_7SP_BRAM_Opt", "RV32IM_8SP_withoutOpt", "RV32IM_8SP"],
         "RV64": ["RV64I_5SP", "RV64IM_5SP", "RV64IM_6SP", "RV64IM_7SP", "RV64IM_7SP_BRAM",
                  "RV64IM_7SP_BRAM_Opt", "RV64IM_8SP_withoutOpt", "RV64IM_8SP"]}
FIELD = {"dhrystone": ("dhrystones_per_sec", "dmips_per_mhz", "Dhrystones/Sec", "DMIPS/MHz"),
         "coremark":  ("iterations_per_sec", "coremarks_per_mhz", "Iterations/Sec", "CoreMarks/MHz")}

HEAD = PatternFill("solid", fgColor="404040")
STALE = PatternFill("solid", fgColor="FFEB9C")
thin = Side(style="thin", color="BFBFBF"); BORDER = Border(thin, thin, thin, thin)

wb = openpyxl.Workbook(); wb.remove(wb.active)
carried = blanked = 0
for fam in ("RV32", "RV64"):
    ws = wb.create_sheet(fam)
    ws["A1"] = f"RV-IM100 — {fam} — FPGA-measured benchmark results"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("One row per bitstream. A yellow blank means the earlier measurement was taken "
                "at a different clock and no longer describes this bitstream — re-measure it.")
    ws["A2"].font = Font(italic=True, size=9); ws.merge_cells("A2:G2")
    cols = [("variant", 24), ("benchmark", 11), ("bitstream", 42), ("clock MHz", 11),
            ("throughput", 15), ("per-MHz score", 15), ("note", 40)]
    for c, (n, _) in enumerate(cols, 1):
        cell = ws.cell(row=4, column=c, value=n)
        cell.fill = HEAD; cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", wrap_text=True); cell.border = BORDER
    row = 5
    for v in ORDER[fam]:
        for b in ("dhrystone", "coremark"):
            m = man.get((v, b))
            if not m:
                continue
            clk = float(m["clock_mhz"])
            thr_f, per_f, thr_h, per_h = FIELD[b]
            old = fb.get(v, {})
            prev = old.get(per_f) or ""
            mm = re.search(r"@\s*([\d.]+)\s*MHz", prev)
            valid = bool(mm) and abs(float(mm.group(1)) - clk) < 0.6
            vals = [v, b, m["bitstream"], round(clk, 3),
                    (int(old[thr_f]) if valid and old.get(thr_f, "").isdigit() else None),
                    (prev if valid else None),
                    ("" if valid else
                     f"re-measure: earlier value was at {mm.group(1)} MHz" if mm else "re-measure")]
            for c, val in enumerate(vals, 1):
                cell = ws.cell(row=row, column=c, value=val); cell.border = BORDER
            if not valid:
                ws.cell(row=row, column=5).fill = STALE
                ws.cell(row=row, column=6).fill = STALE
                blanked += 1
            else:
                carried += 1
            row += 1
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    ws.cell(row=row + 1, column=1,
            value="Throughput is Dhrystones/Sec or Iterations/Sec; per-MHz score is DMIPS/MHz "
                  "or CoreMarks/MHz. FPGA measurements only — tool results are in "
                  "Research_data_complete_*.xlsx.").font = Font(italic=True)
wb.save(DST)
print(f"wrote {DST}")
print(f"  carried over (clock still matches): {carried}/32")
print(f"  blanked for re-measurement        : {blanked}/32")
