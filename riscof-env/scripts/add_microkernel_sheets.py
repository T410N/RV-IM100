#!/usr/bin/env python3
"""Add the micro-kernel campaign to the master workbook.

Three sheets: the derived penalty matrices the paper cites, the 280 matched
pairs those matrices are computed from, and the full 1,526-run table as bulk
evidence.  Contents is extended in place so the existing manual edits to the
workbook survive.
"""
import csv, pathlib, datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

DST = pathlib.Path.home()/"Documents"/"Research_data_MASTER_0921.xlsx"
MK  = pathlib.Path("/home/khwl/Desktop/RV-IM100/riscof-env/microkernels/out")
F   = "NanumGothic"
HDR = PatternFill("solid", fgColor="FFBFBFBF")
WARN= PatternFill("solid", fgColor="FFFFF2CC")

def put(ws, r, c, v, *, bold=False, fill=None, align="left", wrap=False, fmt=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = Font(name=F, size=11, bold=bold)
    if fill: cell.fill = fill
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fmt: cell.number_format = fmt
    return cell

def note(ws, r, text, width=8):
    put(ws, r, 1, text, wrap=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
    ws.row_dimensions[r].height = 30
    return r+1

def order(v):
    """RV32 block then RV64, shallow to deep within each."""
    depth = {"5SP":0,"6SP":1,"7SP":2,"7SP_BRAM":3,"7SP_BRAM_Opt":4,"8SP":5,"8SP_withoutOpt":6}
    name = v.replace("socs_","")
    xlen = 0 if name.startswith("RV32") else 1
    ism  = 0 if "IM_" not in name else 1
    tail = name.split("_",1)[1] if "_" in name else ""
    return (xlen, ism, depth.get(tail, 9), name)

# ---------------------------------------------------------------- read data
runs  = list(csv.DictReader(open(MK/"results.csv")))
pairs = list(csv.DictReader(open(MK/"dependency_penalties.csv")))
variants = sorted({r["variant"] for r in runs}, key=order)
short = lambda v: v.replace("socs_","")

# derived tables, transcribed from ANALYSIS.md
DEP = {  # variant -> (alu g0, alu g1, load g0, load g1, load g2)
 "RV32IM_5SP":(0,0,0,0,0), "RV32IM_6SP":(0,0,1,0,0), "RV32IM_7SP":(0,0,1,0,0),
 "RV32IM_7SP_BRAM":(0,0,1,0,0), "RV32IM_7SP_BRAM_Opt":(0,0,1,0,0),
 "RV32IM_8SP":(1,0,2,1,0), "RV32IM_8SP_withoutOpt":(1,0,2,1,0), "RV32I_5SP":(0,0,0,0,0),
 "RV64IM_5SP":(0,0,0,0,0), "RV64IM_6SP":(0,0,1,0,0), "RV64IM_7SP":(0,0,1,0,0),
 "RV64IM_7SP_BRAM":(0,0,1,0,0), "RV64IM_7SP_BRAM_Opt":(0,0,1,0,0),
 "RV64IM_8SP":(1,0,2,1,0), "RV64IM_8SP_withoutOpt":(1,0,2,1,0), "RV64I_5SP":(0,0,0,0,0)}
BR = {   # variant -> (taken-branch excess, misses, JAL ratio, JALR ratio, matched JAL)
 "RV32IM_5SP":(0.0,1,2.0002,2.0002,2.0), "RV32IM_6SP":(0.0,1,2.0003,2.0003,2.0),
 "RV32IM_7SP":(2.0,1,4.0628,4.0628,4.0), "RV32IM_7SP_BRAM":(2.0,1,4.0628,4.0628,4.0),
 "RV32IM_7SP_BRAM_Opt":(2.0,1,5.0628,5.0628,5.0),
 "RV32IM_8SP":(2.0154,1,6.0941,7.0941,6.0), "RV32IM_8SP_withoutOpt":(2.0154,1,5.0941,6.0941,5.0),
 "RV32I_5SP":(0.0,1,2.0002,2.0002,2.0),
 "RV64IM_5SP":(0.0,1,2.0002,2.0002,2.0), "RV64IM_6SP":(0.0,1,2.0003,2.0003,2.0),
 "RV64IM_7SP":(2.0,1,4.0628,4.0628,4.0), "RV64IM_7SP_BRAM":(2.0,1,4.0628,4.0628,4.0),
 "RV64IM_7SP_BRAM_Opt":(2.0,1,5.0628,5.0628,5.0),
 "RV64IM_8SP":(2.0154,1,6.0941,7.0941,6.0), "RV64IM_8SP_withoutOpt":(2.0154,1,5.0941,6.0941,5.0),
 "RV64I_5SP":(0.0,1,2.0002,2.0002,2.0)}
AR = {   # variant -> (MUL, DIV, DIVW, SW)
 "RV32IM_5SP":(4,36,None,0), "RV32IM_6SP":(4,36,None,0), "RV32IM_7SP":(4,36,None,0),
 "RV32IM_7SP_BRAM":(4,36,None,1), "RV32IM_7SP_BRAM_Opt":(4,36,None,1),
 "RV32IM_8SP":(4,36,None,1), "RV32IM_8SP_withoutOpt":(4,36,None,1),
 "RV32I_5SP":(None,None,None,0),
 "RV64IM_5SP":(4,68,36,0), "RV64IM_6SP":(4,68,36,0), "RV64IM_7SP":(4,68,36,0),
 "RV64IM_7SP_BRAM":(4,68,36,1), "RV64IM_7SP_BRAM_Opt":(4,68,36,1),
 "RV64IM_8SP":(4,68,36,1), "RV64IM_8SP_withoutOpt":(4,68,36,1),
 "RV64I_5SP":(None,None,None,0)}

wb = openpyxl.load_workbook(DST)
for s in ("Microkernel penalties","Microkernel pairs","Microkernel runs"):
    if s in wb.sheetnames: wb.remove(wb[s])

# ------------------------------------------------- sheet 1: derived penalties
ws = wb.create_sheet("Microkernel penalties")
ws.column_dimensions["A"].width = 30
for col in "BCDEFGH": ws.column_dimensions[col].width = 15
r = 1
put(ws,r,1,"Micro-kernel dependency penalties",bold=True); r += 2
r = note(ws,r,"Retirement-bounded ROI delimited by marker instructions.  Each penalty is the cycle "
              "difference between a dependent kernel and its matched independent control with equal "
              "retired-instruction and branch-miss counts, divided by the number of dependencies.  "
              "All 1,526 runs PASS against an independent RV I/M interpreter oracle (rolling ROI hash).")
r += 1

put(ws,r,1,"Matched-pair dependency penalty (extra cycles per dependency)",bold=True); r += 1
hdrs = ["variant","ALU gap 0","ALU gap 1","load gap 0","load gap 1","load gap 2"]
for i,h in enumerate(hdrs): put(ws,r,i+1,h,bold=True,fill=HDR,align="center" if i else "left")
r += 1
for v in variants:
    n = short(v); put(ws,r,1,n)
    for i,x in enumerate(DEP[n]): put(ws,r,i+2,x,align="center")
    r += 1
r += 1

put(ws,r,1,"Control transfer",bold=True); r += 1
hdrs = ["variant","taken-branch excess cycles per retired branch","taken-branch misses",
        "JAL excess cycles per jump (ratio)","JALR excess cycles per jump (ratio)",
        "JAL matched-pair cost"]
for i,h in enumerate(hdrs):
    put(ws,r,i+1,h,bold=True,fill=HDR,align="center" if i else "left",wrap=True)
ws.row_dimensions[r].height = 44; r += 1
for v in variants:
    n = short(v); put(ws,r,1,n)
    for i,x in enumerate(BR[n]):
        put(ws,r,i+2,x,align="center",fmt="0.0000" if isinstance(x,float) else None)
    r += 1
r = note(ws,r,"Only the last column is a matched pair.  The two ratio columns include loop overhead "
              "and pipeline recovery and are descriptive, not isolated misprediction or refill penalties.")
r += 1

put(ws,r,1,"Arithmetic and store service cost (extra cycles per instruction)",bold=True); r += 1
for i,h in enumerate(["variant","MUL","DIV","DIVW","SW"]):
    put(ws,r,i+1,h,bold=True,fill=HDR,align="center" if i else "left")
r += 1
for v in variants:
    n = short(v); put(ws,r,1,n)
    for i,x in enumerate(AR[n]): put(ws,r,i+2,"—" if x is None else x,align="center")
    r += 1
r = note(ws,r,"A dash means the extension is not implemented on that variant.  Operands are the fixed "
              "nontrivial values documented in the generated assembly; M-extension corner cases "
              "(divide by zero, most-negative / -1) are separate correctness kernels, all passing.")
r += 1

put(ws,r,1,"Caveats that must travel with these numbers",bold=True,fill=WARN)
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6); r += 1
for t in ["Penalties characterise the tested instruction pairs.  They do not establish a penalty for "
          "every opcode or dependency type.",
          "Pipeline-occupancy buckets in the run table are a priority partition of observed cycles.  "
          "They are occupancy, not additive counterfactual CPI penalties, and must not be summed into "
          "a stall budget.",
          "RV32IM_8SP shows a HIGHER JAL / JALR cost than RV32IM_8SP_withoutOpt (6 vs 5, and 7 vs 6 "
          "cycles); the same holds on RV64.  The timing-optimised variant is worse on control transfer.  "
          "This is measured and reproducible, but needs an RTL explanation before publication.",
          "Simulation only.  The monitor uses $fopen/$fwrite and is not synthesisable.  The warrant for "
          "simulation fidelity is the Embench FPGA-vs-simulation sheet: 40 of 40 RV64 rows match cycle "
          "and instret exactly."]:
    r = note(ws,r,t,width=6)

# ------------------------------------------------------ sheet 2: matched pairs
ws = wb.create_sheet("Microkernel pairs")
cols = list(pairs[0].keys())
put(ws,1,1,"Matched dependency pairs — the 280 measurements behind the penalty matrix",bold=True)
for i,h in enumerate(cols):
    put(ws,3,i+1,h,bold=True,fill=HDR,align="center")
    ws.column_dimensions[get_column_letter(i+1)].width = max(12,min(26,len(h)+3))
for j,row in enumerate(pairs):
    for i,h in enumerate(cols):
        v = row[h]
        try: v = int(v)
        except ValueError:
            try: v = float(v)
            except ValueError: pass
        put(ws,4+j,i+1,v,align="left" if h in ("variant","kernel","status") else "center")
ws.freeze_panes = "A4"
ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{3+len(pairs)}"

# --------------------------------------------------------- sheet 3: full runs
ws = wb.create_sheet("Microkernel runs")
cols = list(runs[0].keys())
put(ws,1,1,f"Full micro-kernel campaign — {len(runs)} runs, {len(variants)} variants, all PASS",bold=True)
for i,h in enumerate(cols):
    put(ws,3,i+1,h,bold=True,fill=HDR,align="center")
    ws.column_dimensions[get_column_letter(i+1)].width = (
        26 if h in ("kernel","variant") else 18 if "hash" in h or "sha" in h else max(10,len(h)+2))
for j,row in enumerate(runs):
    for i,h in enumerate(cols):
        v = row[h]
        if h not in ("binary_sha256","roi_hash","check_hash"):
            try: v = int(v)
            except ValueError:
                try: v = float(v)
                except ValueError: pass
        c = put(ws,4+j,i+1,v,align="left" if h in ("variant","kernel","family","status","errors") else "center")
        if h == "cpi" and isinstance(v,float): c.number_format = "0.0000"
ws.freeze_panes = "C4"
ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{3+len(runs)}"

# ------------------------------------------------------------ contents update
toc = wb["Contents"]
toc.insert_rows(18, 3)
for i,(name,desc) in enumerate([
  ("Microkernel penalties",
   "Measured dependency, control-transfer and M-extension penalties per variant.  Matched-pair method; "
   "the tables the paper cites."),
  ("Microkernel pairs",
   "The 280 matched dependent/independent pairs the penalty matrix is computed from.  All 280 VALID."),
  ("Microkernel runs",
   f"Every one of the {len(runs)} micro-kernel runs with cycles, retired, event counters, occupancy "
   "buckets and oracle hashes.  All PASS."),
]):
    put(toc,18+i,2,name,bold=True)
    put(toc,18+i,3,desc,wrap=True)
    toc.row_dimensions[18+i].height = 28

gap = next(r for r in range(1,toc.max_row+1) if toc.cell(r,2).value == "Known gaps")
last = toc.max_row + 1
put(toc,last,3,"Micro-kernel results are simulation only; the monitor is not synthesisable.  "
               "Simulation fidelity is established by the Embench FPGA-vs-simulation sheet (40/40 exact).",
    wrap=True)
toc.row_dimensions[last].height = 26

wb.save(DST)
print(f"  wrote {DST}")
print(f"  sheets ({len(wb.sheetnames)}): {', '.join(wb.sheetnames)}")
print(f"  penalties: 16 variants x 3 tables | pairs: {len(pairs)} | runs: {len(runs)} x {len(cols)}")
