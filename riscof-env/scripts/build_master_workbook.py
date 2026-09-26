#!/usr/bin/env python3
"""~/Documents/Research_data_MASTER_0921.xlsx

One workbook holding every measurement made for the revision.  Sheets are
imported from the workbooks that already hold verified data, rather than
regenerated, so nothing is re-derived and possibly changed in transit:

  from 0917_Research_data_FPGA_Complete.xlsx   SoC + core synthesis, CoreMark and
                                               Dhrystone FPGA throughput
  from Research_data_Embench_0918-A.xlsx       Embench FPGA (RV64; RV32 pending)
  from Research_data_Final_0921.xlsx           profiling, Embench simulation,
                                               core comparison, SAIF detail,
                                               verification

The FPGA throughput figures exist only in the workbook -- logs/fpga_benchmarks.csv
is stale, still carrying pre-correction clocks -- so they are copied from the
sheet and not rebuilt from CSV.
"""
import pathlib, datetime, copy
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

H   = pathlib.Path.home()/"Documents"
DST = H/"Research_data_MASTER_0921.xlsx"
F   = "NanumGothic"
HDR = PatternFill("solid", fgColor="FFBFBFBF")

SOURCES = [
    (H/"0917_Research_data_FPGA_Complete.xlsx",
     {"Sheet1":"RV64 SoC and FPGA", "Sheet2":"RV32 SoC and FPGA"}),
    (H/"Research_data_Embench_0918-A.xlsx",
     {"RV64":"Embench FPGA RV64", "RV32":"Embench FPGA RV32",
      "Validation":"Embench FPGA vs simulation"}),
    (H/"Research_data_Final_0921.xlsx",
     {"CPI and stalls":"CPI and stalls",
      "Embench simulation":"Embench simulation",
      "Core comparison":"Core comparison",
      "SAIF detail":"SAIF detail",
      "Normalized metrics":"Normalized metrics",
      "Clocking resources":"Clocking resources",
      "Verification":"Verification"}),
]

def copy_sheet(src, dst):
    """Copy values, formulas, formatting, widths and merges."""
    for row in src.iter_rows():
        for c in row:
            if c.value is None and not c.has_style:
                continue
            d = dst.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                d.font          = copy.copy(c.font)
                d.fill          = copy.copy(c.fill)
                d.alignment     = copy.copy(c.alignment)
                d.border        = copy.copy(c.border)
                d.number_format = c.number_format
    for k, dim in src.column_dimensions.items():
        if dim.width: dst.column_dimensions[k].width = dim.width
    for k, dim in src.row_dimensions.items():
        if dim.height: dst.row_dimensions[k].height = dim.height
    for m in src.merged_cells.ranges:
        dst.merge_cells(str(m))

wb = openpyxl.Workbook(); wb.remove(wb.active)

# ---- contents page, written first so it is the sheet that opens ----
toc = wb.create_sheet("Contents")
toc.column_dimensions["B"].width = 34
toc.column_dimensions["C"].width = 86

imported, missing = [], []
for path, mapping in SOURCES:
    if not path.exists():
        missing.append(path.name); continue
    src = openpyxl.load_workbook(path)
    for sname, target in mapping.items():
        if sname not in src.sheetnames:
            missing.append(f"{path.name}:{sname}"); continue
        copy_sheet(src[sname], wb.create_sheet(target))
        imported.append((target, path.name))

DESC = {
 "RV64 SoC and FPGA":  "Core-only synthesis, SoC utilisation/timing/power, and CoreMark + Dhrystone FPGA throughput, RV64.",
 "RV32 SoC and FPGA":  "The same for RV32.",
 "Embench FPGA RV64":  "Five Embench benchmarks measured on hardware, with CPI, build clock and derived runtime.",
 "Embench FPGA RV32":  "Deliberately blank.  Those runs predate the RV32 counter-read fix; bitstreams are rebuilt and awaiting re-measurement.",
 "Embench FPGA vs simulation": "Every FPGA cycle and instret against RTL simulation.  40 of 40 RV64 rows match exactly.",
 "CPI and stalls":     "CPI, IPC, branch frequency, misprediction rate and an 11-counter stall breakdown, both benchmarks x 16 variants.",
 "Embench simulation": "Full 19 x 16 Embench matrix.  The five also run on FPGA are marked.",
 "Core comparison":    "PicoRV32, RVCoreP and VexRiscv against the RV-IM100 cores under identical synthesis methodology.",
 "SAIF detail":        "Per-component activity-based power for all 32 SoC builds.",
 "Normalized metrics":  "CoreMark score, CoreMark/MHz, CoreMark/LUT, DMIPS/LUT, CoreMark/W and energy per iteration.",
 "Clocking resources":  "BUFGCTRL and PLLE2_ADV per build.  Every design uses a PLL, not an MMCM.",
 "Verification":       "RISCOF, riscv-tests, Embench and randomized (AAPG) pass counts per variant, with the two randomized failure modes explained.",
}

def put(ws, cell, v, *, bold=False, fill=None, align="left", wrap=False):
    c = ws[cell]; c.value = v
    c.font = Font(name=F, size=11, bold=bold)
    if fill: c.fill = fill
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    return c

put(toc,"B2",f"RV-IM100 — complete measurement set", bold=True)
put(toc,"B3",f"built {datetime.date.today().isoformat()}")
put(toc,"B5","sheet",bold=True,fill=HDR); put(toc,"C5","contents",bold=True,fill=HDR)
r = 6
for target, origin in imported:
    put(toc,f"B{r}",target,bold=True)
    put(toc,f"C{r}",DESC.get(target,""),wrap=True)
    toc.row_dimensions[r].height = 28
    r += 1
r += 1
put(toc,f"B{r}","Known gaps",bold=True,fill=HDR); r += 1
for t in ["Embench RV32 on FPGA — bitstreams rebuilt after the counter fix, board runs outstanding.",
          "No SAIF capture exists for the Embench images; the power sheets describe the CoreMark and Dhrystone SoC builds.",
          "Randomized testing shows a control-flow divergence on the 7- and 8-stage variants that no compliance "
          "suite or benchmark reaches.  See the Verification sheet.  It affects no reported measurement.",
          "External soft-core CoreMark/MHz and DMIPS/MHz are not measured here and must be cited from the "
          "respective papers; only area and frequency were measured under matched conditions."]:
    put(toc,f"C{r}",t,wrap=True); toc.row_dimensions[r].height = 26; r += 1
if missing:
    r += 1; put(toc,f"B{r}","Not found at build time",bold=True,fill=HDR)
    for m in missing: r += 1; put(toc,f"C{r}",m)

wb._sheets.insert(0, wb._sheets.pop(wb.sheetnames.index("Contents")))
wb.save(DST)
print(f"  wrote {DST}")
print(f"  sheets ({len(wb.sheetnames)}): {wb.sheetnames}")
if missing: print(f"  MISSING: {missing}")
