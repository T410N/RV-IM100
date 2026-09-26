#!/usr/bin/env python3
"""Fill 'Embench FPGA RV32' in the master workbook from the board results.

FPGA values: bitstream/embench/RV32/Done/<variant>/embench_results.csv (40 rows).
sim cycles:  RTL simulation of the same .mem images, run for this check.
build MHz:   exact PLL output 100*M/(D*O) of the implemented Embench project,
             not the 3-dp XDC readback (57.501 -> 57.5, 71.999 -> 72, 113.999 -> 114).
Layout mirrors 'Embench FPGA RV64' exactly.
"""
import csv, re, pathlib, sys
import openpyxl
from openpyxl.styles import Font, Alignment

ROOT = pathlib.Path("/home/khwl/Desktop/RV-IM100")
DST  = pathlib.Path.home()/"Documents"/"Research_data_MASTER_0921.xlsx"
SIM  = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else None   # dir of <variant>__<bench>.uart
T3   = ROOT/"revision_0922"/"T3"/"soc_rows_from_vivado_reports.csv"
SHEET = "Embench FPGA RV32"

MHZ = {"RV32I_5SP": 45.0, "RV32IM_5SP": 43.0, "RV32IM_6SP": 50.0, "RV32IM_7SP": 57.5,
       "RV32IM_7SP_BRAM": 72.0, "RV32IM_7SP_BRAM_Opt": 92.0,
       "RV32IM_8SP_withoutOpt": 114.0, "RV32IM_8SP": 122.0}
ORDER = ["RV32I_5SP", "RV32IM_5SP", "RV32IM_6SP", "RV32IM_7SP", "RV32IM_7SP_BRAM",
         "RV32IM_7SP_BRAM_Opt", "RV32IM_8SP_withoutOpt", "RV32IM_8SP"]

fpga = {}
for v in ORDER:
    p = ROOT/"bitstream/embench/RV32/Done"/v/"embench_results.csv"
    for r in csv.DictReader(open(p)):
        fpga[(r["architecture"], r["benchmark"])] = (int(r["cycles"]), int(r["instret"]), r["verify"])
assert len(fpga) == 40, len(fpga)

sim = {}
if SIM:
    for f in SIM.glob("*.uart"):
        v, b = f.stem.split("__")
        m = re.search(r"EMBENCH (\S+) cycles=(\d+) instret=(\d+) verify=(\w+)",
                      f.read_text(errors="replace").replace("\x00", ""))
        if m and m.group(1) == b: sim[(v, b)] = (int(m.group(2)), int(m.group(3)), m.group(4))

wb = openpyxl.load_workbook(DST); ws = wb[SHEET]
blocks = {}
for r in range(1, ws.max_row+1):
    h = ws.cell(r, 2).value
    if isinstance(h, str) and h.startswith("RV32  ") and "SAIF" not in h:
        blocks[h.replace("RV32  ", "").strip()] = r
assert len(blocks) == 5, blocks

def put(r, c, v, fmt=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(name="Calibri", size=11)
    x.alignment = Alignment(horizontal="right" if c > 2 else "left", vertical="center")
    if fmt: x.number_format = fmt
    return x

filled = mismatch = 0
for bench, hdr in blocks.items():
    for i, v in enumerate(ORDER):
        row = hdr + 1 + i
        assert ws.cell(row, 2).value == v.replace("RV32", "", 1), (row, ws.cell(row, 2).value, v)
        cyc, ins, ver = fpga[(v, bench)]
        s = sim.get((v, bench))
        put(row, 3, cyc); put(row, 4, ins); put(row, 5, round(cyc/ins, 4), "0.0000")
        if s:
            put(row, 6, s[0])
            exact = (s[0] == cyc and s[1] == ins and s[2] == ver)
            put(row, 7, "EXACT" if exact else "DIFFERS")
            mismatch += 0 if exact else 1
        else:
            put(row, 6, ""); put(row, 7, "not simulated")
        put(row, 8, ver)
        put(row, 9, MHZ[v], "0.######")
        put(row, 10, round(cyc/MHZ[v]/1000.0, 3), "0.000")
        filled += 1

# SAIF block, mirroring RV64: SoC Fmax from the Vivado reports, power from 'SAIF detail' (coremark)
saif_hdr = next((r for r in range(1, ws.max_row+1)
                 if isinstance(ws.cell(r, 2).value, str) and "SAIF" in ws.cell(r, 2).value), None)
if saif_hdr:
    sd = wb["SAIF detail"]
    start = next(r for r in range(1, sd.max_row+1) if sd.cell(r, 2).value == "coremark")
    power = {}
    for r in range(start+1, start+20):
        n = sd.cell(r, 2).value
        if isinstance(n, str) and n.startswith("RV32"):
            power[n] = [sd.cell(r, c).value for c in range(3, 15)]
    fmax = {}
    for r in csv.DictReader(open(T3)):
        if r["bench"] == "coremark" and r["variant"].startswith("RV32"):
            fmax[r["variant"]] = float(r["SoC_Fmax"])
    for i, v in enumerate(ORDER):
        row = saif_hdr + 1 + i
        if ws.cell(row, 2).value is None: continue
        put(row, 3, fmax.get(v, ""), "0.000")
        for j, val in enumerate(power.get(v, [])[:11]):
            put(row, 4+j, val)
        filled += 1

wb.save(DST)
print(f"  filled {filled} rows in '{SHEET}'   sim mismatches: {mismatch}")
